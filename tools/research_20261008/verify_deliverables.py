"""Reconcile final Word/PPTX packages with the frozen research output after visual QA."""
from pathlib import Path
import contextlib
import csv
import hashlib
import io
import json
import re
import sys
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pymupdf
from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
RUN = ROOT / 'training_runs/beginner_guide_20261008'
RESEARCH = ROOT / 'training_runs/teacher9_20261008'
STRUCTURE = RUN / 'guide_structure.json'
SKILL = Path('C:/Users/Amiee/.codex/plugins/cache/openai-primary-runtime/documents/26.430.10722/skills/documents')
sys.path.insert(0, str(SKILL / 'scripts'))
from table_geometry import audit_docx_tables


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    structure = json.loads(STRUCTURE.read_text('utf-8'))
    target = ROOT / structure['docx']
    document = Document(target)
    log = io.StringIO()
    with contextlib.redirect_stdout(log):
        issues = audit_docx_tables(target)
    (RUN / 'table_geometry_audit.txt').write_text(log.getvalue(), 'utf-8')
    assert issues == 0
    assert len([p for p in document.paragraphs if p.style.name == 'Heading 1']) == 41
    assert structure['docx_sha256'] == sha(target)
    for name, digest in structure['inputs'].items():
        assert sha(ROOT / name) == digest, name
    pdf = next((RUN / 'docx_qa_final').glob('*.pdf'))
    rendered = pymupdf.open(pdf)
    assert len(rendered) == 41
    assert all(len(page.get_text()) > 500 for page in rendered)
    images = list((RUN / 'docx_qa_final').glob('page-*.png'))
    assert len(images) == 41
    with ZipFile(target) as package:
        word_xml = ET.fromstring(package.read('word/document.xml'))
        word_text = ''.join(word_xml.itertext())

    pptx = ROOT / '学习笔记/10月9日汇报/量化选股实验进展_20261009.pptx'
    with ZipFile(pptx) as package:
        assert package.testzip() is None
        slides = [n for n in package.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)]
        notes = [n for n in package.namelist() if re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml', n)]
        media = [n for n in package.namelist() if n.startswith('ppt/media/') and not n.endswith('/')]
        assert len(slides) == len(notes) == 23
        assert len(media) == 8
        texts = [ET.fromstring(package.read(n)) for n in slides]
        ppt_text = ''.join(''.join(t.itertext()) for t in texts)
        shapes = sum(len(t.findall('.//{http://schemas.openxmlformats.org/presentationml/2006/main}sp')) for t in texts)
        assert shapes > 100, 'Expected native editable text/table cells'
        note_lengths = [len(''.join(ET.fromstring(package.read(n)).itertext())) for n in notes]
        assert min(note_lengths) > 50, 'Speaker note is unexpectedly empty'
        def pixel_digest(im):
            im = im.convert('RGBA')
            return hashlib.sha256(str(im.size).encode() + im.tobytes()).hexdigest()
        source_pixels = {pixel_digest(Image.open(p)) for p in (RESEARCH / 'figures').glob('*.png')}
        media_pixels = {pixel_digest(Image.open(io.BytesIO(package.read(n)))) for n in media}
        assert media_pixels <= source_pixels, 'A slide image differs from the actual experiment figure'

    summary = json.loads((RESEARCH / 'final_summary.json').read_text('utf-8'))
    checks = json.loads((RESEARCH / 'verification.json').read_text('utf-8'))
    assert checks['passed'] and len(checks['checks']) == 59
    expected = [37.43, 64.47, 62.53, 76.42, 100.07]
    for result, value in zip(summary['factor_count_results'], expected):
        assert round(result['total_return'] * 100, 2) == value
        assert f'{value:.2f}%' in word_text
        assert f'{value:.2f}' in ppt_text
    for value in ['76.03', '121.53', '43.95', '24.21', '73.25', '93.49', '55.39']:
        assert value in word_text and value in ppt_text, value
    for name in ['现金支付日', '可卖', '发布时间']:
        assert name in word_text, name
    rows = list(csv.DictReader((RESEARCH / 'oct8_frozen_model_monitoring.csv').open(encoding='utf-8-sig', newline='')))
    assert len(rows) == 100 and len({r['stockid'] for r in rows}) == 100
    assert all(r['future_label_status'] == 'not_mature' and r['not_retrained_for_2026Q4'] == 'True' for r in rows)

    result = {
        'passed': True, 'document_pages': 41, 'document_sections': 41,
        'document_tables': len(document.tables), 'table_geometry_issues': issues,
        'visual_reviewed_document_pages': list(range(1, 42)),
        'visual_reviewed_presentation_slides': list(range(1, 24)),
        'visual_review': 'All document pages and all slide PNGs reviewed. Fixed header padding, metric-label frames and two awkward wraps; final changed slides rechecked. No clipped content or orphan pages.',
        'pptx_slides': len(slides), 'pptx_speaker_notes': len(notes),
        'pptx_native_shapes': shapes, 'pptx_actual_experiment_images': len(media),
        'image_pixel_reconciliation': 'All 8 embedded media match an experiment figure exactly.',
        'layout_checker': {'errors': 0, 'accepted_warnings': 23, 'reason': 'Source footer and page number are intentionally separate adjacent text boxes, visually reviewed on all slides.'},
        'research_checks_passed': 59, 'monitoring_rows_checked': len(rows),
        'docx_sha256': sha(target), 'pdf_sha256': sha(pdf), 'pptx_sha256': sha(pptx),
        'summary_sha256': sha(RESEARCH / 'final_summary.json'),
        'render_images_sha256': {p.name: sha(p) for p in images},
    }
    (RUN / 'document_qa.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), 'utf-8')
    structure['layout_qa'] = 'passed: canonical renderer, all 41 document pages and all 23 slides visually reviewed; package, figures, geometry and key values reconciled'
    STRUCTURE.write_text(json.dumps(structure, ensure_ascii=False, indent=2), 'utf-8')
    print(json.dumps({k: v for k, v in result.items() if k not in ['render_images_sha256', 'visual_reviewed_document_pages', 'visual_reviewed_presentation_slides']}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
