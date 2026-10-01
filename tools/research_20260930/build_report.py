"""Build the Chinese DOCX with the managed document runtime; render separately for QA."""
from pathlib import Path
import csv,json,sys
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[2]
RUN=ROOT/'training_runs/teacher9_20260930'
OUT=ROOT/'学习笔记/老师意见实施_20260930';OUT.mkdir(parents=True,exist_ok=True)
SKILL=Path('C:/Users/Amiee/.codex/plugins/cache/openai-primary-runtime/documents/26.430.10722/skills/documents')
sys.path.insert(0,str(SKILL/'scripts'))
from table_geometry import apply_table_geometry,column_widths_from_weights
S=json.loads((RUN/'final_summary.json').read_text(encoding='utf-8'))
V=json.loads((RUN/'output_verification.json').read_text(encoding='utf-8'))
A=S['bj']['accounts'];B=S['shsz']['accounts']
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.45)

def font(style,size,bold=False,color='202B33'):
    st=doc.styles[style];st.font.name='Arial';st.font.size=Pt(size);st.font.bold=bold;st.font.italic=False;st.font.color.rgb=RGBColor.from_string(color)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    pf=st.paragraph_format;pf.line_spacing=1.08;pf.space_before=Pt(0);pf.space_after=Pt(6)
for name,size,bold,color in [('Normal',12,False,'202B33'),('Title',22,True,'22577A'),('Subtitle',12,False,'647582'),
    ('Heading 1',16,True,'22577A'),('Heading 2',14,True,'22577A'),('Heading 3',12,True,'22577A'),('Caption',10,False,'647582')]:font(name,size,bold,color)
for name,before,after in [('Heading 1',12,6),('Heading 2',10,4),('Heading 3',8,3)]:
    st=doc.styles[name];st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after)
    st.paragraph_format.keep_with_next=True
header=sec.header.paragraphs[0];header.text='研究实施记录  /  老师九项意见  /  2026-09-30'
header.runs[0].font.size=Pt(9);header.runs[0].font.color.rgb=RGBColor.from_string('647582')
border=OxmlElement('w:pBdr');bottom=OxmlElement('w:bottom')
for k,v in [('val','single'),('sz','4'),('color','D9E1E6'),('space','6')]:bottom.set(qn('w:'+k),v)
border.append(bottom);header._p.get_or_add_pPr().append(border)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('研究用途  ·  ');run.font.size=Pt(9);run.font.color.rgb=RGBColor.from_string('647582')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text,style=None):return doc.add_paragraph(text,style)
def h(text):doc.add_heading(text,1)
def page(title):
    heading=doc.add_heading(title,1);heading.paragraph_format.page_break_before=True
def pct(value):return f'{value:+.2%}'
def f(value):return f'{value:.3f}'
def table(headers,rows,weights,size=10.5):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
    for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
    for values in rows:
        for cell,text in zip(t.add_row().cells,values):cell.text=str(text)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(0)
                para.paragraph_format.line_spacing=1.06
                if ri==0:para.paragraph_format.keep_with_next=True
                for run in para.runs:run.font.size=Pt(size);run.font.bold=ri==0
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F1F5F7');cell._tc.get_or_add_tcPr().append(shade)
    apply_table_geometry(t,column_widths_from_weights(weights,9360),table_width_dxa=9360,indent_dxa=120)
    return t
def figure(name,caption,width=6.5):
    para=doc.add_paragraph();para.paragraph_format.space_after=Pt(2)
    para.paragraph_format.keep_with_next=True
    para.add_run().add_picture(str(RUN/'figures'/(name+'.png')),width=Inches(width))
    p(caption,'Caption')
def links(items):
    para=p('','Caption')
    for i,(label,url) in enumerate(items):
        if i:para.add_run('  ·  ')
        link=OxmlElement('w:hyperlink');link.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True))
        run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'22577A');props.append(color)
        run.append(props);txt=OxmlElement('w:t');txt.text=label;run.append(txt);link.append(run);para._p.append(link)
def metrics_rows(metrics,names):
    return [[label,pct(metrics[name]['total_return']),pct(metrics[name]['max_drawdown']),f(metrics[name]['sharpe'])] for name,label in names]

# Page 1: conclusions and decision implications, not a decorative cover.
doc.add_heading('老师九项意见\n实施结果与效果报告',0)
p('2026 年 9 月 30 日  ·  独立研究版本 teacher9_20260930','Subtitle')
p('现有文件能够支撑的研究工作已经实际运行：股票池、因子审计与筛选、滚动训练、现金账户、指数对照、HMM 及分钟执行测试均有代码和可核查产物。原先的首轮实验保留，未覆盖。')
table(['本轮范围','实际完成'],[
    ['北交所','8 个季度滚动；2024-10-08—2026-09-29 账户评价'],
    ['沪深样本','31 个季度滚动；2019-01-02—2026-08-21 账户评价'],
    ['验证','12 项功能测试；16 条账户、39 轮训练边界检查通过']],[1.3,5.2],11)
doc.add_heading('效果：有改进，也暴露了明确问题',1)
p(f"北交所筛选因子 LightGBM 的研究账户累计 {pct(A['pred_selected']['total_return'])}，价格模型为 {pct(A['pred_price']['total_return'])}。但主模型最大回撤 {pct(A['pred_selected']['max_drawdown'])}，2026 年截至 9 月 29 日为 −26.60%，不能据累计收益判断策略稳定。")
p(f"Ridge 对照累计 {pct(A['pred_ridge']['total_return'])}，优于 LightGBM；HMM 的收益和夏普均不如固定 85% 仓位上限。沪深样本筛选模型累计 {pct(B['pred_selected']['total_return'])}，没有证明跨市场有效。")
p('分钟压力回放只完成约 67.5% 的目标股数；因子额外延迟一天，主模型累计收益降至 +27.27%。执行与数据可用时间对结果影响很大。')
doc.add_heading('这些数字能说明什么',1)
p('上述收益是扣除假设成本的未复权价格研究账户结果。真实分红送转台账、完整历史股票资格与商业平台执行尚未验收，因此不能视为可直接实现的投资收益。本轮完成的是可复现研究与问题定位，尚未完成实盘策略验收。')

# Page 2: one-to-one mapping to the teacher's nine questions.
page('九项意见逐项对照')
p('“已实现”表示代码和本轮产物已经运行；不代表所有数据口径和策略有效性均通过正式验收。')
table(['问题','本次实际工作与效果','当前边界'],[
 ['1  动态 Top100','按过去 20 日成交额每日重建；最多 100 只；保留未入池原因与因子缺失记录。','实现完成；资格仍用历史有效交易记录代理。'],
 ['2  资金曲线与基准','现金/股数记账，成本、T+1、延期卖出；统一指数和模型对照。','研究账户完成；真实公司行为待补。'],
 ['3  2500+ 字段与去重','2508 字段全匹配字典；2447 为数值候选。训练内日截面相关去重。','全量完成最新快照审计；历史建模覆盖 292 个 Alpha 字段。'],
 ['4  单因子和模型比较','每日 Rank IC、HAC、BY-FDR、五分组；LightGBM/Ridge/单因子对照。','本轮完成；Ridge 更好，未据此回调测试参数。'],
 ['5  沪深长历史','300 只可用文件样本，独立 31 轮价格特征实验。','完成试验；非全市场，含幸存者偏差。'],
 ['6  真正滚动训练','每季重新筛选、验证早停、重训，清除跨边界标签。','北交所 8 轮、沪深 31 轮完成。'],
 ['7  HMM 解释和对照','两状态、月度重估、前向过滤；输出概率、转移和仓位图。','未通过增量效果判断；收敛诊断有异常。'],
 ['8  时点、复权与单位','增加一天延迟测试、复权差异诊断、日分钟单位核对、历史股本面板。','已推进；真实发布时间与公司行为仍缺。'],
 ['9  vn.py 与执行','分钟订单回放、版本化信号、时间门控、离线 BarData 适配器。','准备完成；商业版环境未安装验收。']],[1.15,3.4,1.95],10.5)

# Page 3: what was downloaded and how the universe changed.
page('数据与股票池：先确定每天能研究谁')
p('本轮只读下载所需文件，保存来源清单及 SHA-256。北交所 raw/qfq 各 350 个日线文件；沪深确定性抽样 300 只，各取 raw/qfq；另接入 23 个基准文件、120 表最新快照、Alpha 历史、61 个分钟文件和 071/601 历史股本。')
figure('universe','图 1  股票池按日重建，成交量与成交额得到的名单并不相同。',6.3)
p('入池规则：至少 60 个有效交易日，最近 20 个市场交易日中至少 16 日有有效交易，当日价格、成交量和成交额有效；再按过去 20 日平均成交额降序取前 100。早期不足 100 只时保留真实数量，不强行凑足。')
p(f"北交所累计纳入 301 只股票、110110 条股票日记录；非空股票池为 75—100 只，{S['universe']['days_below100']} 日不足 100。成交量池与成交额池的重合率中位数为 79%。入池后平均 Alpha 字段覆盖率为 94.49%；11946 条记录低于 70%，仍保留缺失信息。")
p('这修复了静态名单和“因子缺失就换股票”的问题。上市、退市、ST、停牌等正式历史资料仍不完整；当前文件集合也可能遗漏已消失的股票，历史有效交易记录不能替代完整股票身份表。')

# Page 4: historical factor testing and rolling process.
page('因子筛选与滚动：只让过去参与决策')
p('供应商字典已匹配全部 2508 个非键字段，其中 2447 个数值候选、61 个名称/来源/日期辅助字段。最新快照中，数值候选有 9 个常数字段、22 个覆盖率低于 70%；这些是当日质量结论，不是全历史有效性结论。')
p('历史主实验使用 Alpha101、Alpha191 共 292 个字段，加 12 个价格与流动性特征，共 304 个候选。每轮仅用训练段计算每日横截面 Spearman IC，按覆盖率与 IC 天数筛选，绝对 IC 前 100 名进入去重，日截面相关均值绝对值达到 0.85 则剔除，最多保留 40 个。')
figure('factor_selection','图 2  2026Q3 训练段选中因子的 IC 与相关矩阵（节选 18 个）。',6.3)
p('每轮用过去 504 个市场日期作训练段、126 个作验证段；因五日标签跨期，实际训练尾部作清除。验证确定早停轮数后，用测试季度之前的可用标签重训；外层测试季度不参与选因子或调参。39 轮边界检查全部通过。')
p('每轮保留因子卡片、筛除原因、相关矩阵、模型及时间边界。HAC 使用 5 阶滞后，BY-FDR 处理多重比较；q 值作为证据报告，不作为本次额外调参条件。最后一轮选出 7 个 Alpha101、30 个 Alpha191 和 3 个价格特征。')

# Page 5: reference account comparison.
page('北交所：账户有改善，回撤仍然很大')
figure('bj_equity','图 3  同一时间区间、相同账户规则的对照。曲线含假设成本，未计真实公司行为。',6.3)
table(['模型 / 对照','累计收益','最大回撤','夏普'],metrics_rows(A,[
 ('pred_selected','筛选因子 LightGBM'),('pred_price','价格 LightGBM'),('pred_ridge','Ridge'),
 ('pred_simple','训练选定单因子'),('pool_equal','Top100 等权目标')])+[['北证 50 价格指数',pct(A['pred_selected']['benchmark_price_return']),'—','—']],[2.7,1.3,1.3,1.2])
p('账户起始 100 万元，每日新批次预算最多为当日净值 20%，选前 5 只；单股上限 10%，买入额受过去 20 日平均成交额的 1% 约束。t+1 开盘买入、t+5 收盘卖出，不能卖出则延期；现金不为负。末五日停止新开批次，以评价完整持有期。')
p('单边合计成本为 15bp 假设；涨跌停使用保守价格阈值代理，未覆盖所有特殊规则。Top100 对照实行等权目标，但受最小股数等约束，平均仓位仅 46.68%，不能当作无摩擦满仓等权指数。')

# Page 6: statistical signal and sensitivity; no misleading compounded quintiles.
page('效果稳健性：方向、延迟和成本都要看')
figure('quintiles','图 4  外层测试五分组均值。五日标签有重叠，这不是复利资金曲线。',6.1)
p(f"主模型 479 个可评价日期的平均 Rank IC 为 {S['bj']['prediction']['pred_selected']['mean_ic']:.4f}；10 日区块、2000 次自助法的均值 95% 区间为 [0.1255, 0.1837]。相对价格模型的 IC 增量区间为 [0.0029, 0.0240]，支持本样本内存在排序改善；不能由此推断可成交利润。")
figure('cost_random','图 5  随机账户中位数 −12.08%，范围 −39.89%—+9.03%，均低于主模型；仅 10 组，不构成正式检验。',6.0)
p('单边成本 5/15/30bp 时，主模型累计分别为 +93.58% / +64.47% / +28.79%。仅给 Alpha 特征额外滞后一天并重训后为 +27.27%，IC 为 0.1291；这说明历史因子真实可用时点必须核实。')

# Page 7: honest extension scope.
page('沪深长历史：流程跑通，结果不支持推广')
p('从当前可用文件中，按股票代码 SHA-256 排序，沪市 60 开头与深市 00 开头各取 150 只。该样本不是历史指数成分，也不是全部沪深股票；带有可用文件与幸存者偏差。')
figure('shsz_equity','图 6  2019-01-02—2026-08-21 的独立长历史对照。“300 只样本”不代表沪深 300 成分股。',6.1)
table(['模型 / 对照','累计收益','最大回撤','夏普'],metrics_rows(B,[
 ('pred_price','价格 LightGBM'),('pred_selected','筛选价格特征 LightGBM'),('pred_ridge','Ridge'),('pool_equal','Top100 等权目标')])+
 [['沪深 300 价格指数',pct(B['pred_selected']['benchmark_price_return']),'—','—']],[2.7,1.3,1.3,1.2])
p('沪深实验只使用 12 个价格/流动性候选特征，已完成 31 季度的筛选、验证和重训；未扩展全部 292 个 Alpha 历史，不能声称已完成完整跨市场因子验证。')
p('筛选模型平均 Rank IC 为 0.0592，账户却亏损；相对价格模型的 IC 增量区间为 [−0.0052, 0.0007]，未呈现明确改善。排序相关性、交易成本、公司行为与持仓收益是不同问题，不能用正 IC 代替账户验收。')

# Page 8: state interpretation, controls and diagnostics.
page('HMM：状态更可解释，但未带来增量效果')
figure('hmm','图 7  概率是下一交易日预测；仓位上限在信号日之后应用。实际仓位通常低于上限。',6.05)
table(['仓位方案','累计收益','最大回撤','夏普'],metrics_rows(A,[
 ('hmm_cap','两状态 HMM'),('fixed85_cap','固定 85% 上限'),('vol_cap','简单波动率上限')]),[2.7,1.3,1.3,1.2])
p('观察变量是当日池内平均收益、20 日市场波动对数、上涨下跌广度和离散度；每月仅用此前 504 日标准化并重估，按训练期波动均值标记低/高状态。使用前向过滤，低/高状态对应 100%/70% 上限，再做 0.8/0.2 平滑。')
p('24 次拟合的最短预期状态持续期约 12.89 日；但 10 次入选拟合的最后一步似然略降，库的 converged=True 不能当作严格收敛证明。已另存严格诊断，重算后的状态概率与原输出一致。HMM 夏普低于固定上限约 0.167，故暂不采用为有效增强。')

# Page 9: data semantics.
page('价格、股本与时点：修正口径，保留未决项')
doc.add_heading('复权文件不能按名称直接代入收益公式',2)
p('本次北交所检查窗口未发现非正 qfq 收盘价，但 raw 与 qfq 的关系主要呈加法调整特征。47872 对外层测试标签中，492 对的 qfq 比率收益与 raw 价格收益相差超过 1 个百分点；没有负价并不证明比率收益正确。')
p('按 raw/qfq 仿射关系变化标记出 3963 个候选调整边界，756 条测试标签跨越候选边界。剔除这些标签后主模型 IC 为 0.1517，仅是诊断性敏感度结果；边界识别用到未来诊断信息，不能当作可交易过滤条件。')
doc.add_heading('历史股本已整理，但没有冒充真实公司行为',2)
table(['项目','本次结果'],[
 ['历史快照','071/601 两表，共 1184 个日期、350 只股票、259656 行'],
 ['池内字段覆盖','流通股本约 99.983%；自由流通股本 100%'],
 ['口径冲突','22954 行出现自由流通股本 > 流通股本，或流通股本 > 总股本；已标记，异常行不生成换手率代理'],
 ['建模使用','作为辅助候选保存，未事后加入已评估主模型；未从股本变化推断分红送转']],[1.45,5.05],11)
p('股本按每日文件日期关联，不以报告期日期回填到历史。来源含报告期口径，可能滞后或修订；文件日 18:00 可用只是研究假设，仍需供应商确认真实发布时点。')
doc.add_heading('单位与记账',2)
p('分钟成交量汇总 / 日线成交量中位数为 99.665，支持日线“手”、分钟“股”的约 100 倍换算；金额汇总比中位数为 0.99666，说明两者仍有覆盖差异。换手率代理采用日线量 ×100 / 历史流通股本。')
p('账户引擎已有显式现金分红和股数调整接口，并通过模拟守恒测试；本次真实账户未输入经核实的分红送转台账。接口正确不等于真实公司行为已经处理完毕。')

# Page 10: execution and verification.
page('分钟执行与 vn.py：交付可核查的离线准备')
p('检查 61 个分钟文件：2025Q4 的 60 个日期，加 2026-09-24 单位样本。重复键检查通过；空价格、零量和指数行不能当作可成交股票行情。')
table(['2025Q4 订单级压力回放','结果'],[
 ['日线参考订单 / 有分钟文件覆盖','558 / 558 笔'],
 ['请求股数 / 回放成交股数','2033110 / 1372966 股'],
 ['股数完成率','67.53%'],
 ['全量成交订单比例','43.37%'],
 ['完全未成交订单比例','0%；其余多为部分成交'],
 ['相对日线参考价的不利价差中位数','约 10.52bp；单笔值可为负，不等同总组合成本']],[3.4,3.1],11)
p('入场使用 09:31—10:00，收盘卖出使用 14:50—15:00。预先存在的订单按完成分钟的最不利价格、实际成交量 1% 回放，同股同分钟共享容量。容量用于事后压力测试，不能在分钟开始前预知。')
p('该回放只测试日线已成交参考订单，没有重新引入被拒订单，未重算独立现金资金曲线，也没有队列、冲击或全套交易规则。67.53% 的股数完成率说明日线参考收益不能直接移植为分钟可实现收益。')
doc.add_heading('离线交接与验证',2)
p('已导出 2420 条带模型版本、可用时点、最早执行日的研究信号；最新信号日为 9 月 29 日，按假设最早用于 9 月 30 日。vnpy_adapter.py 提供 BJ→BSE 映射、时间门控和分钟 BarData 转换；商业版环境未安装、未连接、未发单，尚不能声称平台迁移完成。')
p(f"12 项测试通过，覆盖记账、T+1、延期、零量、限价代理、公司行为模拟及防未来信息。{V['actual_account_files_checked']} 条账户的现金加市值误差最大 {V['max_cash_plus_value_error']:.2g}；39 轮边界与去重检查通过，首轮 16 个文件哈希不变。")

# Page 11: concrete reproducibility and outstanding acceptance work.
page('交付位置、复现方式与剩余验收')
p('下表 R 表示 training_runs/teacher9_20260930/。所有路径相对工作区根目录。','Caption')
table(['交付内容','工作区相对路径'],[
 ['本轮研究代码 / 操作说明','tools/research_20260930/README.md'],
 ['冻结输入与来源清单','training_data/research_20260930/'],
 ['统一结果与验证','R/final_summary.json、output_verification.json'],
 ['每季模型、筛选证据','R/bj/、R/shsz/；每季含模型、因子卡片与边界'],
 ['资金、成交、分钟回放','R/*_account.csv、*_trades.csv、\nminute_order_replay.csv、execution_signals.csv'],
 ['学习笔记入口','学习笔记/README.md，原方案顶部已添加实施说明']],[1.7,4.8],10.5)
p('复现顺序：冻结输入 → prepare → 字典映射 → train → accounts → replay → prepare_capital → test_research → verify_outputs → analyze → build_report。命令见代码目录 README；研究使用 .venv_quant，Word 与字典使用文档运行时。')
doc.add_heading('剩余工作对应明确验收条件',2)
p('数据：补正式历史股票资格，核对分红送转日期、现金和股数；确认因子与股本的发布/修订时点，处理 22954 条股本口径冲突。')
p('执行：接入真实公司行为、费用和交易规则并对账；分钟回放需联合管理现金、持仓、未成交订单与容量，随后在交付的 vn.py 环境验收。')
p('研究：保留本次正负结果，用新时间段验证模型，扩展沪深 Alpha 历史和股票范围；不能根据本次测试收益调参后仍称独立验证。')
doc.add_heading('口径参考与证据',2)
p('数值来自本地冻结输入和验证产物。日线限价检查为研究代理，未完整复刻下列正式规则；点击名称可查原文。','Caption')
links([('北交所交易规则','https://www.bse.cn/jygl_list/200028217.html'),
       ('北交所历史规则','https://www.bse.cn/jygl_list/200010919.html'),
       ('vn.py 组合策略文档','https://www.vnpy.com/docs/cn/community/app/portfolio_strategy.html'),
       ('vn.py 交易所枚举','https://github.com/vnpy/vnpy/blob/master/vnpy/trader/constant.py')])

doc.core_properties.title='老师九项意见：实施结果与效果报告'
doc.core_properties.subject='2026-09-30 冻结研究、结果与验收边界'
doc.core_properties.author='量化研究项目'
dest=OUT/'老师九项意见_实施结果与效果报告.docx';doc.save(dest)
# Searchable companion notes are local; only the requested Word is delivered to the user.
lines=[]
for element in doc.element.body:
    if element.tag==qn('w:p'):
        text=''.join(element.itertext())
        # Use python-docx paragraph objects instead of XML itertext's duplicate node content.
        from docx.text.paragraph import Paragraph
        para=Paragraph(element,doc);text=para.text
        if text:lines.append(text)
    elif element.tag==qn('w:tbl'):
        from docx.table import Table
        t=Table(element,doc)
        for row in t.rows:lines.append(' | '.join(cell.text.replace('\n',' ') for cell in row.cells))
(OUT/'实施结果_可检索文本.md').write_text('\n\n'.join(lines),encoding='utf-8')
print(dest)
