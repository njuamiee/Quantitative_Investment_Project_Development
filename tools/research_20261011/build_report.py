"""Build Chinese project documentation with the managed document runtime."""
from pathlib import Path
import json,csv,sys,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[2];RUN=ROOT/'training_runs/research_20261011'
OUT=ROOT/'学习笔记/20261012_多期限选股项目';OUT.mkdir(parents=True,exist_ok=True)
if '--preview-pending' in sys.argv:
    OUT=RUN/'document_draft';OUT.mkdir(exist_ok=True)
SKILL=Path('C:/Users/Amiee/.codex/plugins/cache/openai-primary-runtime/documents/26.430.10722/skills/documents')
sys.path.insert(0,str(SKILL/'scripts'));from table_geometry import apply_table_geometry,column_widths_from_weights

def load(p):return json.loads((RUN/p).read_text(encoding='utf-8'))
def readcsv(p):return list(csv.DictReader((RUN/p).open(encoding='utf-8-sig',newline='')))
def pct(x):return f'{float(x):.2%}'
def num(x):return '—' if x is None else f'{float(x):.2f}'
H=[1,2,5,10,20];M={h:load(f'h{h}/candidate_metrics.json') for h in H};V={h:load(f'h{h}/main_period_metrics.json') for h in H}
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1);sec.header_distance=sec.footer_distance=Inches(.42)
for name,size,bold,color in [('Normal',12,False,'263742'),('Title',25,True,'174E70'),('Subtitle',12,False,'607585'),('Heading 1',17,True,'174E70'),('Heading 2',13,True,'174E70'),('Caption',9,False,'607585')]:
    s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor.from_string(color);s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    s.paragraph_format.line_spacing=1.08;s.paragraph_format.space_after=Pt(6)
    if name.startswith('Heading'):s.paragraph_format.keep_with_next=True
header=sec.header.paragraphs[0];header.text='多期限选股研究  /  2026-10-09数据截止  /  2026-10-12交付';header.runs[0].font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.add_run('研究候选 · 连续账户 · 可复核   |   ').font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
transcript=[]

def p(text,style=None):
    transcript.append(text);return doc.add_paragraph(text,style)
def page(title):
    h=doc.add_heading(title,1);h.paragraph_format.page_break_before=True;transcript.append('\n## '+title+'\n')
def sub(title):doc.add_heading(title,2);transcript.append('\n### '+title+'\n')
def table(headers,rows,weights=None,size=9.5):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
    for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
    widths=column_widths_from_weights(weights or [1]*len(headers),9360)
    apply_table_geometry(t,widths)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1);para.paragraph_format.line_spacing=1.06
                for r in para.runs:r.font.size=Pt(size);r.font.bold=ri==0
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7F0F5');cell._tc.get_or_add_tcPr().append(sh)
    transcript.append('| '+' | '.join(headers)+' |');transcript.append('| '+' | '.join(['---']*len(headers))+' |')
    for row in rows:transcript.append('| '+' | '.join(str(x).replace('\n','；') for x in row)+' |')
    spacer=p('','Caption');spacer.paragraph_format.space_after=Pt(2);spacer.paragraph_format.line_spacing=Pt(2);return t
def fig(name,caption,width=6.4):doc.add_picture(str(RUN/'figures'/name),width=Inches(width));p(caption,'Caption')
def link(label,url):
    para=doc.add_paragraph();rel=para.part.relate_to(url,RT.HYPERLINK,is_external=True);h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'176B91');pr.append(c);r.append(pr);tx=OxmlElement('w:t');tx.text=label;r.append(tx);h.append(r);para._p.append(h);transcript.append(f'[{label}]({url})')

p('多期限选股项目\n完整研究说明与实验报告','Title')
p('10月9日会议要求落实｜SFTP更新复核｜HMM｜日内因子｜LLM与Jev','Subtitle')
p('输入快照：2026-10-11；行情截止：2026-10-09。交付日期：2026-10-12。用户已确认飞书无需继续核对。')
sub('本轮得到了什么')
p('完成1、2、5、10、20交易日训练，形成五个完整历史回测同时达到“夏普>1、年化Calmar>2”的探索候选。但5日和20日候选最近一年仍亏损，尚未完成稳定达标模型的验收。')
table(['预测 / 持有','年化收益','夏普¹','Calmar','最大回撤'],[[f"{h}日 / {M[h].get('holding_period',h)}日",pct(M[h]['annual_return']),num(M[h]['sharpe_rf15']),num(M[h]['calmar']),pct(M[h]['max_drawdown'])] for h in H],[1.35,1,1,1,1])
p('¹ 夏普采用年化1.5%无风险收益；账户现金收益记0。完整区间为2024-10-08至2026-10-09，共487个交易日；所有结果扣除每边5bp假定成本。','Caption')
sub('必须和成绩一起看的条件')
p('这些候选经过历史比较后筛选，属于回顾性研究；完全按预定验证规则自动选型的方案没有达到Calmar>2。旧历史也曾在前几轮研究中查看，因此不能称为全新留出集通过。')
p('5日和20日候选收益主要来自早期阶段。最近一年、最近半年均亏损；15bp成本与信号延迟压力通过，也不能弥补近期失效。短窗口月度更新的进一步检验见第08.1节。')
p('较小回撤依赖低仓位。五个候选平均实际持仓约6%—9%；10日与20日最终候选均预测对应远期收益、实际采用5日持有。真实数据发布版本、退市完整性和分红到账时间仍有缺口。')

page('01  数据更新、能做的实验与缺口')
A=load('audit_summary.json')
p(f"本轮只读盘点SFTP：120张因子表已更新至10月9日；下载并留存{A['downloaded_manifest_files']}个文件，失败{A['failed_downloads']}个。与10月8日本地快照比较，{A['revision_counts']['changed']}个同名文件哈希变化、{A['revision_counts']['new_to_snapshot']}个新纳入、{A['revision_counts']['same']}个相同。哈希变化不自动等于历史数值被修订，也可能只是新增尾部日期。")
table(['资料','已核实状态','本轮用途'],[
 ['北交所日线/复权/基准','新快照到10月9日；350只价格来源','多期限标签、日线成交和基准'],
 ['120张逐日因子表','目录最新日期10月9日','更新c01/c02等表，335项候选输入'],
 ['北交所分钟与清洗Tick','1184个日期文件，最新9月30日','125个本地分钟文件的IC/执行检验'],
 ['minute-all.zip','中央目录1184项，也只有BJ','排除压缩包内另藏SH/SZ分钟数据'],
 ['沪深旧日线目录','三个复权目录各5868项；样本内容到8月21日','复用已有496只样本长期检验'],
 ['factor_lib','抽样Alpha001、0_1到9月30日','不把因子最新日误当行情最新日']],[1.15,1.85,1.7])
p('交易所公告确认2026年10月8日恢复交易。因此10月8、9日分钟文件属于有效交易日缺口。原始Tick压缩包文件名可到10月2日，但该日休市，未把它当作新增可交易数据。')
p(f"逐个日期/路径的已识别缺口共{A['date_gaps']}项，见 missing_files.csv。该清单按各目录首次可用日期起与日历比对，包含文件日期缺失及不可用目录；不代表每个缺口对本轮模型都有影响。")
p('优先补：BJ 10月8/9日分钟数据、SH/SZ分钟目录和8月21日后的活跃样本行情、历史证券主表/退市/ST/涨跌停字段、因子发布时点与历史版本、分红实际到账及送转可卖日、带版本时间的财报新闻与人工标签。完整字段要求见 missing_fields_and_documents.csv。')

page('02  项目从输入到明日持仓如何运行')
table(['环节','明确规则'],[
 ['数据快照','原文件+远端元信息+SHA256；保留此前实验，不覆盖成同一口径'],
 ['每日股票池','过去20日平均成交额Top100；至少60个有效历史日，近20日至少16日有效成交'],
 ['收盘特征','同日截面排名转为[-0.5,0.5]附近；缺值填中性0；供应商发布时间仍是假设'],
 ['季度模型','训练504交易日、验证126日；按标签终点剔除跨边界样本；季度开始前重训'],
 ['月度状态','过去504日拟合HMM；当日收盘递推滤波，输出次日状态概率和仓位上限'],
 ['次日执行','用前一收盘信号；未复权价格成交；现金、股数、成本与延期卖出连续记录'],
 ['验收','保留全部固定方案、自动选择、事后优化结果，分别报告']],[1.05,4])
sub('预测标签与持有周期')
table(['预测期限','标签定义','本轮最终候选持有'],[[str(h)+'日','t+1开盘→t+2开盘' if h==1 else f't+1开盘→t+{h}收盘',str(M[h].get('holding_period',h))+'日'] for h in H],[1,2.6,1.4])
p('1日期限采用跨交易日开盘收益，避免A股当日买入又卖出。其他标签使用后复权价格消除拆分/分红造成的机械跳变；账本仍按未复权价格成交。20日预测、5日执行表示模型学习20日相对收益，组合每批5日退出并用更新信号再配置；并未把5日标签改名为20日标签。')
p('股票池每天更新；因子筛选和监督模型每季度更新；HMM每月更新。这三种频率不同，任何一项变动都不能重置总资金。')

page('03  因子：来源、含义与筛选数学')
p('输入共335项：供应商Alpha101/191共292项、19项本地价量特征、24项LLM提出的可执行算术表达式。factor_catalog.csv逐项列出来源、公式或原始字段、说明和可用时点假设。编号本身不表示经济效应。')
table(['例子','公式 / 含义','训练中如何用'],[
 ['gtalpha046','(MA3+MA6+MA12+MA24)/(4×收盘价)：均线相对当前价','正IC常对应近期回落的反转信号'],
 ['gtalpha063','6期平滑上涨幅度/绝对变动幅度×100','RSI类短期强弱；可出现负IC'],
 ['gtalpha067','上述比率改为24期平滑','中期强弱；与其他因子去重'],
 ['gtalpha071','(收盘价-MA24)/MA24×100','价格相对均线偏离'],
 ['LLM风险调整动量','mom20/(vol20+0.001)','假说为风险调整趋势；需要实测']],[1.1,2.45,1.55])
p('每个训练日先计算因子截面排名与未来收益截面排名的Pearson相关，即Spearman RankIC。对训练日逐日IC取等权均值；只接受覆盖率≥70%的候选，按平均IC绝对值排序，再剔除与已入选项训练内相关绝对值≥0.92的冗余项。最终模型使用前60或100项。')
p('负IC不直接丢弃：IC组合使用带符号的平均IC权重，并归一化绝对权重之和；Ridge和树模型从训练样本学习方向。供应商因子做同日截面排名，但原始文件历史发布时点未验证，所以另做整个评分延迟一日的压力检验。')
p('例：2025Q4的5日训练卡片中，gtalpha046平均IC约+0.145，gtalpha063约−0.138，gtalpha067约−0.137。数字只来自该折训练，不能当作未来固定有效性。所有季度卡片保留，避免只展示效果好的因子。')

page('04  模型训练、验证与选择')
table(['模型','本轮设定及作用'],[
 ['Ridge60','60项、L2强正则，线性可解释基线'],['LightGBM100','100项、浅树、小学习率、叶子样本与L2限制'],
 ['LambdaRank60','收益截面分档，按交易日分组学习排序'],['CatBoost60 / XGBoost60','两类提升树，检验对模型实现的依赖'],
 ['ExtraTrees60','随机树集成、限制深度和叶子样本'],['LLM-only Ridge','只用24个LLM提案，单独检验有无增益'],
 ['IC / conditional IC','带符号IC组合；三状态概率加权的条件因子权重'],['Ensemble','Ridge、LightGBM、CatBoost、XGBoost截面排名均值']],[1.5,3.5])
p('每个预测期限覆盖2024Q4—2026Q4共9个季度预测块，最后一块只有10月8/9日。验证先按逐日RankIC保留三个候选模型，再检验Top10/20及15种风险规则；目标为max min(验证夏普、验证Calmar/2)。验证结束后用最近504日已成熟标签重训。最后几日未成熟标签不参与IC，仍可以产生交易评分并在期末计入未平仓市值。')
table(['期限','预定验证选择夏普','Calmar','最大回撤'],[[str(h)+'日',num(V[h]['sharpe_rf15']),num(V[h]['calmar']),pct(V[h]['max_drawdown'])] for h in H],[1,1.6,1,1.4])
p('上表为截至2026Q3的预定选择方案；没有一个通过本次双门槛。候选成绩更高，来自后续固定方案对照和风险研究，必须保留“事后筛选”标签。共500组固定外层对照、540组风险缩放/保护对照和216组长期预测执行扩展，另有各折验证网格；大量尝试增加选择偏差。')

page('05  资金账本、费用与风险规则')
p('初始100万元。每个交易日按前一日信号新建批次，原始默认新批次不超过净值/持有天数；再受现金、总仓位和单股上限限制。任何季度都不重新发放资金。')
table(['约束','实际实现'],[
 ['T+1','逐批记录入场日，卖出必须晚于入场日；1日策略先在次日开盘平上一批'],
 ['股数与集中度','初始买入≥100股，整数增量；新买入后单股不超过当时净值10%'],
 ['成交容量','单只股票买卖共享过去20日平均成交额1%的日成交金额上限'],
 ['买不全/卖不出','缺报价、停牌代理、涨跌停代理、容量不足均保留失败/部分成交；未卖完继续持仓'],
 ['费用','买卖每边5bp总成本研究假设；另列0/1/2/15bp，不把假设称为券商账单'],
 ['风险上限','固定仓位、12%/20%波动目标、趋势、HMM防守/增长、回撤和CPPI；只读此前信息'],
 ['分红与期末','除权日即时税前现金及送转股可卖情景；最后一天继续正常开仓，未到期持股按市值记账']],[1.1,3.9])
p('年化收益=(期末净值/初始资金)^(365.25/实际日数)−1；夏普=日超额收益均值/日收益标准差×√252；Calmar=年化收益/最大回撤。最大回撤的峰值包含初始100万元。会议举例使用过“累计收益/回撤”，报告另存该比率，验收表统一使用更规范的年化口径。')
p('9项合成测试覆盖T+1、现金恒等式、共用容量、缺失行情、拆股、初始资金回撤、时间因果性、HMM前向滤波、DSL与Jev接口。真实结果另检查5个期限的资金恒等式、无同日卖出、容量、预测唯一性和45折训练标签边界。')

page('06  五个候选的连续净值和资金利用')
fig('candidate_nav_drawdown.png','全部曲线都含现金和费用。历史筛选候选；未声称全新留出集通过。')
table(['期限','评分 / HMM','风险缩放','平均持仓','近空仓日'],[[str(h)+'日',M[h]['model']+' / '+M[h]['base_policy'],num(M[h]['risk_scale']),pct(M[h]['mean_exposure']),pct(M[h]['cash_days_fraction'])] for h in H],[.6,2.25,.65,.8,.8],9)
p('低仓位是回撤下降的重要来源。最低100股约束还会使部分名义Top20无法成交，产生价格与交易粒度偏好；报告保留失败订单数，不能把名义选股数当作实际持股数。将资金扩大十倍也不能机械放大收益，需重新核验股数、容量和成交成本。')

page('07  HMM如何介入，以及确实改善了什么')
p('观测向量包含市场平均收益、对数波动率、趋势/波动比、对数Amihud非流动性、截面收益分散度。训练内标准化并截尾；2/3状态Gaussian HMM，月度使用此前504日拟合，两个固定随机种子按训练似然和占用/持续期诊断选择。92组月度/状态数拟合中1组回退；通过的拟合均未触及220次迭代上限。')
p('αₜ ∝ 发射似然(Oₜ) × (αₜ₋₁ A)；用于下一交易日的是αₜA。每次重训按训练内条件波动排序对齐状态；所有交易端只使用前向结果。防守上限由收缩后的状态均值决定；增长上限由状态收益/波动决定；概率加权后只做过去向指数平滑。')
fig('hmm_controls.png','相同5日IC评分、相同成本。最后一项把固定仓位调到约7%的实际平均敞口，只作为事后解释对照。')
C=load('h5/candidate_hmm_controls.json');match=C[-1]
p(f"5日增长HMM候选夏普{num(M[5]['sharpe_rf15'])}、Calmar{num(M[5]['calmar'])}、回撤{pct(M[5]['max_drawdown'])}；实际敞口匹配的固定仓位分别为{num(match['sharpe_rf15'])}、{num(match['calmar'])}、{pct(match['max_drawdown'])}。这个对照支持本轮HMM的历史择时增益。")
p('增益不是普遍的：未缩放的HMM防守相对匹配仓位提高夏普和Calmar，但最大回撤未必更小；HMM波动率方案也可能损害夏普。状态条件均值—方差因子择时实验并未达标，失败结果保留。')

page('08  状态概率、真实持仓与分段风险')
fig('hmm_probabilities_exposure.png','上图是按波动排序的状态概率，不应直接把“低波”改名为“牛市”。下图显示仓位上限与真实持仓差异。')
segments=load('h5/candidate_subperiods.json')
table(['5日候选分段','年化收益','夏普','Calmar','回撤'],[[{'first_year':'前一年','second_year':'最近一年','recent_half_year':'最近半年'}[x['period']],pct(x['annual_return']),num(x['sharpe_rf15']),num(x['calmar']),pct(x['max_drawdown'])] for x in segments],[1.5,1,1,1,1])
b=load('h5/candidate_bootstrap.json')
p(f"20日块自助法（1500次）得到5日候选夏普95%区间[{num(b['sharpe_rf15_95'][0])}, {num(b['sharpe_rf15_95'][1])}]，Calmar区间[{num(b['calmar_95'][0])}, {num(b['calmar_95'][1])}]。这反映历史样本不确定性，尚未校正全部研究尝试造成的选择偏差；不能把通过比例解释成未来达标概率。")

def short_window_section():
    if not (RUN/'short_windows/summary.json').exists() and '--preview-pending' in sys.argv:
        page('08.1  Short-window results pending in this layout preview')
        return
    page('08.1  发现近期失效后，补做了哪些实验')
    ss=load('short_windows/summary.json');short=readcsv('short_windows/all_results.csv')
    p('先对500组固定方案、540组风险缩放/保护方案重新分段核对，没有方案同时通过完整区间与最近一年或最近半年。随后明确登记新实验，检验“504日训练太旧、季度更新太慢”的假说；这是发现问题后的探索，仍不构成全新留出集。')
    table(['改变什么','固定范围'],[
     ['监督模型','预测1/5/20日；63/126/252日训练窗；每月更新，标签终点剔除；从初始训练期确定的候选因子中再选40项'],
     ['模型复杂度','Ridge正则20/180、深度3的LightGBM、带符号IC；4个固定模型'],
     ['HMM与仓位','504/126日HMM训练窗；三状态防守/增长×0.25/0.5；另有固定与波动对照'],
     ['执行及验收','Top20、每边5bp、相同连续账本；20日预测另对比5/20日持有；三个区间分别要求夏普>1且Calmar>2']],[1.1,3.9])
    p(f"实际完成225个月度训练分组、9份预测文件、{ss['tested']}个账户方案。完整区间通过{ss['full_pass']}个，最近一年通过{ss['recent_year_pass']}个，最近半年通过{ss['recent_half_pass']}个，同时三个区间通过{ss['all_periods_pass']}个。月度预测的标签边界、非空评分和唯一性检查均通过。")
    def worst_short(x):return min(float(x[k])/(2 if k.endswith('calmar') else 1) for k in ['sharpe_rf15','calmar','recent_year_sharpe_rf15','recent_year_calmar','recent_half_sharpe_rf15','recent_half_calmar'])
    best_short=[max((x for x in short if int(x['forecast_horizon'])==h),key=worst_short) for h in [1,5,20]]
    table(['预测/持有','训练窗/模型','完整 S/C','近一年 S/C','近半年 S/C'],[[f"{x['forecast_horizon']}/{x['holding_period']}日",f"{x['window']}日/{x['model']}",num(x['sharpe_rf15'])+'/'+num(x['calmar']),num(x['recent_year_sharpe_rf15'])+'/'+num(x['recent_year_calmar']),num(x['recent_half_sharpe_rf15'])+'/'+num(x['recent_half_calmar'])] for x in best_short],[1,1.3,1,1.1,1.1],9)
    p('S为无风险1.5%夏普，C为年化Calmar。上表每个期限选择三个区间中最差的min(S,C/2)最高者，仅用于解释跟进实验；其风险规则和所有480条结果完整保存在short_windows/all_results.csv。不能把不同方案在不同区间的最好成绩拼接成一条策略。','Caption')
    if ss['all_periods_pass']==0:p('结果不支持“仅缩短窗口并加快更新就能解决失效”。下一轮须优先解决因子发布时间、证券主表与更多独立行情，并研究近期选股效应变化；目前不能把完整历史区间的达标结果作为稳定验收通过。')
    else:p('本轮有探索方案通过三个历史区间；由于这些区间已用于诊断与选择，仍需冻结后前瞻验证。所有方案与失败结果均保留，不能将其改称事前验收通过。')

short_window_section()

page('08.2  为何IC仍为正，账户却会亏损')
p('固定候选评分重新分段后，发现近期相对排序仍有信息，但5日和20日选中股票的平均远期绝对收益已经为负。只做多的相对排名模型，无法仅凭正IC保证每个阶段赚钱。')
signal=readcsv('recent_signal_diagnosis.csv')
table(['期限/区间','平均RankIC','入选股标签收益','股票池标签收益'],[[x['horizon']+'日/'+('前一年' if x['period']=='first_year' else '最近一年'),f"{float(x['mean_daily_rank_ic']):.3f}",pct(x['mean_top_label_return']),pct(x['mean_pool_label_return'])] for x in signal],[1.4,1,1.3,1.3])
p('此表逐日等权平均已成熟的预测标签，标签终点不跨分段截止；无成本和实际成交约束，期限重叠。它用于诊断信号，不是账户收益、独立交易样本或最终20日候选的5日实际持有收益。','Caption')
bench=load('recent_benchmark_diagnosis.json')
table(['最近一年基线','区间累计收益','口径'],[['动态股票池等权',pct(bench[1]['total_return']),'相同5日批次与资金/成交限制'],['北证50',pct(bench[3]['total_return']),'价格指数，未含分红与跟踪成本']],[1.5,1.2,2.3])
p('最近一年5日候选平均实际持仓只有约3.13%，区间累计亏损约1.90%；相对基线显著减少损失，但仍是亏损。HMM在这里主要控制风险敞口，不能把所有缺乏正绝对收益的股票变成盈利交易。')
p('这一诊断支持下轮把“横截面排序”和“是否有足够绝对预期收益覆盖成本”作为两个可分别检验的模型问题。需训练内估计并校准绝对收益/不交易门槛，保留同期现金和简单趋势对照，再在冻结区间评价。不能根据已经看到的下跌日期事后关闭交易，也不能靠放大仓位把负期望变成正期望。')
p('本页说明现象与可检验假说，不把市场下跌视为所有损失的唯一原因。供应商历史版本、幸存者偏差、价格代理与小额整股成交偏差仍需要独立排查。')

page('09  稳健性：哪些候选更值得继续')
table(['5日候选情景','成本/边','夏普','Calmar','最大回撤'],[[['成本0bp','成本2bp','主情景5bp','成本15bp','不计分红','评分延迟1日'][i],str(x['cost_bp'])+'bp',num(x['sharpe_rf15']),num(x['calmar']),pct(x['max_drawdown'])] for i,x in enumerate(load('h5/candidate_sensitivity.json'))],[1.5,1,1,1,1])
table(['期限','15bp仍双通过','延迟1日仍双通过','判断'],[[str(h)+'日','是' if load(f'h{h}/candidate_sensitivity.json')[3]['pass_target'] else '否','是' if load(f'h{h}/candidate_sensitivity.json')[-1]['pass_target'] else '否','近期表现未通过' if h in [5,20] else '门槛较脆弱'] for h in H],[1,1.4,1.5,1.6])
p('完整区间五个候选达到数值门槛，但1/2日期限对成本或信号时点敏感；10日期限的部分压力情景也未通过。5日与20日通过成本和延迟压力，仍未通过近期表现检验。20日预测最终实际持有5日；原20日持有候选Calmar约1.47，保存在h20/aligned_holding目录。')
p('分红未验证真实支付日、税率和送转可卖日，因此主情景是税前即时到账近似；不计分红的压力结果单列。涨跌停也只是30%价格代理，缺ST与历史精确参考价，不能据此宣称可直接实盘复制。')

page('10  日内因子、分钟执行与沪深扩展')
p('125个分钟文件形成12114个“股票—日”聚合样本。分钟IC使用前一日已知股票池，剔除午休与缺失分钟跨越，检验1/5/20分钟未来收益；成交信号使用完整分钟后的信息。分钟预测IC不能直接解释为A股可当天往返的利润。')
p('日聚合特征包括尾盘收益、实现/下行波动、偏度、成交量集中度、尾盘量占比、收盘/VWAP偏离、午后相对上午强弱与价量相关等。2025Q4训练、2026Q3评价；固定Ridge后次日交易。')
ir=load('minute_factor_results.json')
table(['分钟聚合→日预测','风险规则','夏普','Calmar'],[[str(x['horizon'])+'日',x['policy'],num(x['sharpe_rf15']),num(x['calmar'])] for x in ir if x['policy']=='none'],[1.6,1.4,1,1])
ex=load('minute_execution_results.json')
table(['分钟参与率','成交股数比例','最低现金','容量超额/同日卖出'],[[pct(x['participation']),pct(x['share_fill_fraction']),f"{x['minimum_cash']:.0f}",f"{x['capacity_overrun']} / {x['same_day_sales']}"] for x in ex],[1,1.4,1.2,1.8])
p('分钟执行压力：买入9:31—10:00、卖出14:50—15:00，买用分钟高价、卖用低价，0.5/1/2%成交量参与率。此项复用原Top5账本，对新的IC评分按5bp检验，不是最终低仓位候选的分钟净值证明。')
sh=load('shsz_cost5_results.json')
table(['沪深31季度旧样本重跑','累计收益','最大回撤'],[[x['model'],pct(x['total_return']),pct(x['max_drawdown'])] for x in sh],[2,1.5,1.5])
p('沪深保留496只来源股票的旧样本与原账本局限，只把费用改为5bp重跑；不把8月21日后的缺失行情、退市信息和板块交易规则假装已补齐。')

page('11  LLM因子挖掘：已完成的本地闭环')
p('本轮由Codex提出24个明确假说的公式，记录名称、输入、经济解释和信息截止时点。表达式只允许字段、数字及加减乘除，拒绝函数调用、属性访问、索引和外部执行；除零转缺值。因果滚动特征由审查过的代码提供，生成表达式本身不能向未来移动数据。')
table(['步骤','实现与证据'],[
 ['假说','如风险调整动量、短期反转、下行风险、量价压力；llm_factor_proposals.json'],
 ['实现/审核','factor_dsl.py及拒绝执行测试；不把LLM描述直接作为可交易证据'],
 ['训练筛选','仅训练内IC、覆盖与去重；与供应商特征共同进入335项候选'],
 ['独立消融','LLM-only Ridge单列，避免因合并模型盈利就声称LLM增益'],
 ['因子择时','63日因子收益动量；HMM状态条件均值/波动/Sharpe与收缩配置'],
 ['留痕','所有公式、季度因子卡、预测、失败组合、账本和输入哈希保留']],[1.15,3.9])
fixed=readcsv('h5/fixed_comparisons.csv')
table(['5日同风险规则比较','夏普','Calmar'],[[x['model']+' / '+x['policy'],num(x['sharpe_rf15']),num(x['calmar'])] for x in fixed if x['model'] in ['llm_ridge','ic','conditional_ic'] and x['policy']=='hmm3_defensive'],[3,1,1])
p('本地因子提案与执行闭环已经实际运行；未调用RD-Agent收费后端，也没有把手工写入的真实数据标签伪称LLM生成结果。两篇因子择时论文提供研究假说，本轮近似实验未复制原始跨市场样本与论文完整方法。')

page('12  Jev：离线接入完成，真实调用为零')
table(['未来任务','API类型','输出和使用边界'],[
 ['财报信息分类','choice','改善/恶化/混合/证据不足及概率；相同报告期与历史版本比较'],
 ['新闻事件评分','score','五档0—4的概率加权评分；是基本面影响等级，不等同未来收益'],
 ['风险信号提取','noul','是否存在明确的流动性、违约、会计、监管、治理或经营风险证据概率']],[1.4,.7,2.9])
p('模型固定为jev-1.13.0；正式端点为POST /v1/systemone。用户明确要求先不付费使用，因此本轮没有真实API调用，也没有把会议或行情上传到外部模型。默认call_jev拒绝联网；只有后续明确允许后才能更改调用开关并安全配置环境变量。')
p('每条文本保留发行人、文档ID、来源、发布时间、摄取时间、版本可得时间。三个时间必须不晚于决策截止点；之后发布或后来修订的文本不可回填到过去。文档内容被当成待分类证据，不作为程序指令。')
p('响应校验检查问题集合、类型、类别、概率归一化、0—4评分与概率加权一致性、置信度及实际响应模型版本。本地测试显式使用合成样例，不能记成真实财报识别准确率。')
p('接入真实数据后的验收：先做人工标签及时间切分，验证集拟合校准，报告Brier、LogLoss、ECE与分类准确率；最终按相同选股/成交规则比较“有Jev/无Jev”。风险概率只作为辅助特征或风险上限，不直接当作股票上涨概率。')
p('后续仍需真实文档档案与标签、中文任务校准、真实API契约检查，以及对组合的增量检验。代码和离线检验已完成，业务效果尚未实测。')

page('13  前沿模型调研与Kronos本地实测')
table(['路线','本轮采取的做法','下一步条件'],[
 ['Qlib表格/排序模型','实际运行Ridge、提升树、随机树、排序与集成','继续用相同数据和账本比较'],
 ['时序神经网络/图模型','调研Qlib官方模型族，未冒充全部跑完','更充分序列样本、行业关系、算力与独立协议'],
 ['RD-Agent / AlphaGPT','采用假说、实现、反馈和人工审核流程','若启用外部模型需另计成本与资料权限'],
 ['Kronos-mini','公开4.1M模型，本机CPU推理，固定代码/权重提交号','扩大滚动评价，并核实预训练数据时间范围'],
 ['Jev','完成文本分类、评分、风险接口；零真实调用','真实文本标签和后续授权']],[1.2,2,1.8])
kr=load('kronos_results.json')
table(['Kronos预测期限','平均RankIC','截面日期数'],[[str(x['horizon'])+'日',f"{x['mean_rank_ic']:.4f}",x['dates']] for x in kr['results']],[1.8,1.6,1.6])
p('Kronos检验为2026年7—8月的9个约周频截面，共821个股票—日期预测；128根历史日K线输入、单条采样路径、同时提取1/2/5/10/20日输出。20日RankIC较高，但只有9个日期且标签重叠，不能用这个数字宣布稳定收益或独立显著性。')
p('该检验是本机实际推理，未纳入最终低回撤候选的收益冒充来源；预训练样本时间范围也未独立核实。仓库与权重SHA、推理参数和逐股预测已保存。调研覆盖与现有资料匹配的主要路线，不宣称穷尽全球模型。')

tasks=readcsv('requirements_matrix.csv')
for batch in range(4):
    part=tasks[batch*8:(batch+1)*8];page(f'14.{batch+1}  老师和项目要求逐项落实')
    p('原始依据为10月9日两段会议录音文字、本地项目规划、HMM任务PDF及本轮用户要求。源头和详细证据路径完整保存在requirements_matrix.csv。','Caption')
    table(['编号 / 要求','本轮解决方案','状态与剩余条件'],[[x['id']+'\n'+x['requirement'],x['solution'],x['status']] for x in part],[1.2,2.6,1.55],9.5)
    if batch==3:p('没有代替用户向老师或同学发送消息、提交材料或安排会议。相关外部沟通事项由用户自行决定；本次交付已准备可审阅的材料。')

page('15  冻结、复现与下一轮验收')
p('每个期限已导出最终候选参数、最新季度模型、模型/输入哈希和10月9日收盘排名。最新排序对应下一交易时段，属于研究输出，未向券商下单。')
table(['文件/目录','用途'],[
 ['tools/research_20261011','全部新增研究、Jev、测试和文档生成代码'],
 ['training_runs/research_20261011/h*','各折训练记录、固定对照、探索网格、连续账本、候选压力检验'],
 ['h*/frozen_release.json','冻结日期、模型/风险参数、预测与持有期限、哈希'],
 ['h*/frozen_latest_ranking.csv','10月9日收盘评分、排名与风险上限；不是成交保证'],
 ['missing_files.csv / missing_fields_and_documents.csv','逐路径日期缺口和缺失字段/资料'],
 ['requirements_matrix.csv / sources.md','30项要求验收和原始来源'],
 ['artifact_checks.json / verification.json','真实结果检查及9项合成测试']],[2.2,2.8],9)
p('本地复现：powershell -ExecutionPolicy Bypass -File tools/research_20261011/run_local.ps1。默认不联网；-Sync才只读访问SFTP。Jev仍只生成离线样例。Kronos公开权重推理、文档生成与规范渲染分别执行，详见代码README。')
p('前瞻验收建议冻结至少126个交易日：固定参数，仅按已约定的季度模型/月度HMM频率更新；评价期间不看亏损后改风险阈值。若再次优化，另建研究版本。同步补齐发布时点、证券主表和真实成交规则，才能把历史条件性结果升级为更强的可执行证据。')
p('本轮完成的是可重跑的研究模型、真实本地实验、缺口清单与完整说明。未来收益、真实Jev效果、实盘成交和外部数据补齐不能通过修改回测数字提前完成。')

page('15.1  能手算，才能解释模型')
p('以下数值为教学例子，不是实际股票收益。实际日期与评分示例见下一页。')
sub('IC为什么可能是负的')
p('四只股票的因子排名为[1,2,3,4]，未来收益排名为[4,1,2,3]。无并列排名时，RankIC=1−6×Σ(排名差²)/(n×(n²−1))=1−6×12/(4×15)=−0.20。负号表示因子越大，未来相对排名倾向越差；带符号IC组合会反向使用，不是简单删除。真实研究对每个交易日先算截面IC，再平均。')
sub('T+1与费用怎样进入账本')
p('信号在t日收盘形成，最早t+1开盘买。1日模型用t+1开盘→t+2开盘的标签；5日模型用t+1开盘→t+5收盘。若标签终点跨过训练截止点，这条训练样本要剔除。')
p('初始现金1,000,000元，买100股×10元，每边5bp：买入费0.50元，剩余现金998,999.50元。收盘10.20元时，股票市值1,020元，净值1,000,019.50元。下一交易日以10.50元全部卖出，卖出费0.525元，最终现金1,000,048.975元，净赚48.975元。当天买入不能当天卖出；实际还检查成交量、价格及股数限制。')
sub('HMM怎样变成次日风险上限')
p('两状态教学例：昨日后验[0.7,0.3]，转移矩阵两行为[0.9,0.1]和[0.2,0.8]，今日先验=[0.69,0.31]。今日观测在两状态下的似然为[0.8,0.2]，归一化后验≈[0.8990,0.1010]。再乘转移矩阵得到次日概率≈[0.8293,0.1707]。假设训练得到两状态风险上限[1,0.1]，概率加权约0.8464，再乘0.25风险缩放得到约21.16%。实际模型的状态上限由训练内统计确定。')
sub('Calmar为什么不能用两年总收益直接除')
p('假设一年净值从100到120，中途峰值120跌到114：最大回撤5%，年收益20%，Calmar=4。若120是两年后终值，年化收益变为√1.2−1≈9.54%，Calmar≈1.91。相同总收益不能混用不同区间的分母和年化口径。')

page('15.2  真实单日输出与答辩核对')
latest=readcsv('h5/frozen_latest_ranking.csv')
p('以下来自5日历史候选在2026年10月9日收盘后的真实输出。IC是已训练的因子加权评分，不是单只股票上涨概率，也不是未来收益百分比。10月12日尚未有本快照中的成交行情，因而这里只展示决策输入。')
table(['排名','股票代码','IC组合评分','组合风险上限'],[[x['rank'],x['stockid'],f"{float(x['ic']):.4f}",pct(x['target_cap'])] for x in latest[:5]],[.6,1.4,1.4,1.6])
p('该日组合上限约1.65%，低于完整回测平均实际持仓。最终买入需结合已有持仓、可用现金、各批次剩余期限、100股下限和成交容量；表中股票排名高，不等于一定买入，更不等于全仓买入。完整排序（交易规则取前20）和模型哈希保存在frozen_latest_ranking.csv与frozen_release.json。')
table(['老师可能追问','应能清楚说明'],[
 ['每天究竟更新什么？','每天更新池、特征和排名；主要模型季度重新训练，HMM月度重训。短窗口跟进实验另用月度监督模型。'],
 ['HMM有效的证据是什么？','相同预测与相同成本，补平均敞口匹配对照；只支持本历史样本中的择时增益。'],
 ['为什么回撤下降这么多？','平均持仓约6%—9%、现金占比高，HMM与缩放共同作用；并非满仓风险被神奇消除。'],
 ['达到老师要求了吗？','完整历史探索候选达到两个数字，预定选择和近期稳定性验收尚未达到。'],
 ['Jev现在能报准确率吗？','不能；目前只有接口及合成离线测试，还需真实文本、人工标签和校准。'],
 ['下一步最有用的工作是什么？','补发布时间和证券主表，按固定规则积累前瞻数据；把近期失效作为研究问题。']],[1.45,3.55])

page('16  主要来源与引用说明')
for label,url in [
 ('Qlib官方仓库及模型族','https://github.com/microsoft/qlib'),('RD-Agent官方仓库','https://github.com/microsoft/RD-Agent'),
 ('R&D-Agent-Quant原始论文','https://arxiv.org/abs/2505.15155'),('Alpha-GPT 2.0原始论文','https://arxiv.org/abs/2402.09746'),
 ('Kronos官方仓库','https://github.com/shiyu-coder/Kronos'),('Kronos原始论文','https://arxiv.org/abs/2508.02739'),
 ('Jev API契约','https://docs.typesafe.ai/api'),('Jev模型版本说明','https://docs.typesafe.ai/models'),
 ('Factor Momentum and the Momentum Factor，2022','https://doi.org/10.1111/jofi.13131'),
 ('Factor Timing，NBER WP26708，2020','https://www.nber.org/system/files/working_papers/w26708/w26708.pdf'),
 ('北交所2026年休市公告','https://www.bse.cn/important_news/200027428.html')]:link(label,url)
p('本地依据：1009会议记录两段文字、项目规划.mm、任务6 HMM PDF、天软191因子算法手册、供应商registry及旧实验资料。原文件中的连接凭据不复制到报告。完整采用范围和版本说明见sources.md。')
p('报告所有收益来自本地生成的账户结果；外部论文、GitHub项目与供应商文档仅用于方法来源和接口核对。报告没有把第三方宣传成绩作为本项目绩效。')

path=OUT/'多期限选股项目_完整研究报告_20261012.docx';doc.save(path)
(OUT/'项目说明_可检索版.md').write_text('\n\n'.join(transcript),encoding='utf-8')
for file in ['candidate_summary.csv','validation_selected_summary.csv','missing_files.csv','missing_fields_and_documents.csv','requirements_matrix.csv','factor_catalog.csv','sources.md']:
    shutil.copy2(RUN/file,OUT/file)
if (RUN/'short_windows/all_results.csv').exists():shutil.copy2(RUN/'short_windows/all_results.csv',OUT/'短窗口月度更新_全部480组.csv')
brief='''# 十分钟汇报稿

1. 第0—1分钟：目标是每边5bp后夏普>1、年化Calmar>2。完整历史探索有达标候选，但5日和20日最近一年亏损，稳定达标验收尚未完成。
2. 第1—3分钟：展示五期限表、近期分段与短窗口跟进实验。20日预测实际5日持有；低回撤伴随约7%—8%平均持仓，不能只展示全区间好成绩。
3. 第3—5分钟：展示HMM相同IC信号、相同平均敞口的对照。解释只用前向滤波、月度重训和状态软仓位。
4. 第5—6分钟：展示15bp、信号延迟、分红情景。1/2/10日期限的脆弱性也说明，不隐藏失败。
5. 第6—7分钟：日内因子与分钟执行分别汇报；LLM提案已形成受限表达式闭环。Jev仅离线接入，尚无真实业务效果。
6. 第7—8分钟：Kronos本机完成9个截面；RankIC结果只说明值得扩大研究，不是收益保证。
7. 第8—9分钟：列出10月8/9日分钟、沪深新行情、证券主表、发布版本和分红到账等缺口。
8. 第9—10分钟：提出冻结126交易日前瞻观察；请求老师确认费率口径、可接受利用率及剩余数据优先级。

答辩准备：能手算IC符号、1日T+1标签、一次买卖后现金/市值变化、HMM次日概率和Calmar。个人理解仍需本人演练。
'''
(OUT/'十分钟汇报稿.md').write_text(brief,encoding='utf-8');print(path,flush=True)
