"""Check final document structure and evidence integrity after visual review."""
from pathlib import Path
import csv
import hashlib
import json
from collections import Counter
from docx import Document
from docx.oxml.ns import qn
import fitz

root = Path(__file__).resolve().parents[2]
run = root / 'training_runs/teacher_answers_20261012'
out = root / '学习笔记/20261012_老师新问题逐项解答'
dest = next(out.glob('*.docx'))
digest = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
source_hashes = json.loads((out / 'source_hashes.json').read_text('utf-8'))
assert all(digest(root / p) == value for p, value in source_hashes.items())
manifest = json.loads((out / 'delivery_manifest.json').read_text('utf-8'))
assert all(digest(out / p) == value for p, value in manifest.items())
checks = json.loads((out / 'checks.json').read_text('utf-8'))
assert len(checks) == 16 and all(row['passed'] for row in checks)
coverage = json.loads((out / '问题覆盖表.json').read_text('utf-8'))
assert [row['id'] for row in coverage] == [f'Q{i:02}' for i in range(1, 22)]
doc = Document(dest)
assert len(doc.tables) == 20
for table in doc.tables:
    widths = [int(col.get(qn('w:w'))) for col in table._tbl.tblGrid]
    assert sum(widths) == 9360
    for row in table.rows:
        assert [int(cell._tc.tcPr.tcW.get(qn('w:w'))) for cell in row.cells] == widths
qa = run / 'docx_qa_final'
assert {p.name for p in qa.glob('page-*.png')} == {f'page-{i}.png' for i in range(1, 23)}
with fitz.open(next(qa.glob('*.pdf'))) as pdf:
    assert len(pdf) == 22
    page_text_lengths = [len(page.get_text().strip()) for page in pdf]
    assert min(page_text_lengths) > 300
record = {
    'document': str(dest.relative_to(root)),
    'document_sha256': digest(dest),
    'pages': 22,
    'visual_review': 'All 22 final page PNGs inspected during this task; no clipping, overlap, unreadable glyphs, or broken tables found.',
    'table_geometry_passed': 20,
    'question_coverage_count': len(coverage),
    'selection_consistency_checks_passed': len(checks),
    'source_hashes_verified': source_hashes,
    'delivery_manifest_verified': True,
    'pdf_page_text_lengths': page_text_lengths,
    'new_teacher_protocol_retrained': False,
}
(run / 'document_qa.json').write_text(json.dumps(record, ensure_ascii=False, indent=2), 'utf-8')
print(json.dumps({k: v for k, v in record.items() if k != 'source_hashes_verified'}, ensure_ascii=False, indent=2))
