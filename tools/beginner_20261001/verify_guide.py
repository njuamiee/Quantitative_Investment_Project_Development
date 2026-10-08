"""Check document structure and its quoted evidence after all-page visual review."""
from pathlib import Path
import sys,json,hashlib,io,contextlib
import pymupdf
from docx import Document
ROOT=Path(__file__).resolve().parents[2]
RUN=ROOT/'training_runs/beginner_guide_20261002'
S=json.loads((RUN/'guide_structure.json').read_text(encoding='utf-8'))
target=ROOT/S['docx'];doc=Document(target)
SKILL=Path('C:/Users/Amiee/.codex/plugins/cache/openai-primary-runtime/documents/26.430.10722/skills/documents')
sys.path.insert(0,str(SKILL/'scripts'))
from table_geometry import audit_docx_tables
stream=io.StringIO()
with contextlib.redirect_stdout(stream):issues=audit_docx_tables(target)
(RUN/'table_geometry_audit.txt').write_text(stream.getvalue(),encoding='utf-8')
assert issues==0
assert len([p for p in doc.paragraphs if p.style.name=='Heading 1'])==35
pdf=next((RUN/'docx_qa_final').glob('*.pdf'));render=pymupdf.open(pdf)
assert len(render)==35
for i,p in enumerate(render,1):
    assert len(p.get_text())>500,(i,'unexpected short/orphan page')
assert len(list((RUN/'docx_qa_final').glob('page-*.png')))==35
for path,digest in S['inputs'].items():
    assert hashlib.sha256((ROOT/path).read_bytes()).hexdigest()==digest,path
research=json.loads((ROOT/'training_runs/teacher9_20260930/final_summary.json').read_text(encoding='utf-8'))
checks=[('bj','pred_selected',64.47),('bj','pred_price',14.40),('bj','pred_ridge',88.24),
        ('bj','hmm_cap',44.69),('bj','fixed85_cap',65.05),('bj','pred_lag1',27.27),('shsz','pred_selected',-40.35)]
for market,name,expected in checks:
    assert round(research[market]['accounts'][name]['total_return']*100,2)==expected
assert round(research['bj']['accounts']['pred_selected']['max_drawdown']*100,2)==-45.33
audit=json.loads((ROOT/'training_data/audit_20261002/comparison.json').read_text(encoding='utf-8'))
assert audit['market_missing_paths']==39 and audit['factor_missing_paths_2026']==0
assert len(audit['factor_arrivals_since_previous'])==120
old=json.loads((ROOT/'training_runs/teacher9_20260930/frozen_old_run_hashes.json').read_text(encoding='utf-8'))
for name,digest in old.items():
    assert hashlib.sha256((ROOT/'training_runs/lgbm_hmm_20260920'/name).read_bytes()).hexdigest()==digest,name
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
result={'passed':True,'pages':35,'sections':35,'tables':len(doc.tables),'table_geometry_issues':issues,
    'visual_reviewed_pages':list(range(1,36)),
    'visual_review':'All final page PNGs opened and reviewed: text, wrapping, tables, margins, headings, footer; no orphan pages or clipped content.',
    'docx_sha256':sha(target),'pdf_sha256':sha(pdf),
    'old_frozen_files_unchanged':len(old),'key_quoted_returns_verified':len(checks),
    'audit_beijing':audit['audit_end_beijing'],'supplement_beijing':audit['supplement_end_beijing'],
    'research_version':'teacher9_20260930; no new training in this task',
    'render_images_sha256':{p.name:sha(p) for p in sorted((RUN/'docx_qa_final').glob('page-*.png'))}}
(RUN/'document_qa.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
S['layout_qa']='passed: canonical renderer, all 35 final page PNGs visually reviewed; geometry and content checks passed'
(RUN/'guide_structure.json').write_text(json.dumps(S,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k not in ['render_images_sha256','visual_reviewed_pages']},ensure_ascii=False,indent=2))
