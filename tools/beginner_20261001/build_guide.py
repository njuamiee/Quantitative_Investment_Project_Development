"""Create a beginner-first DOCX from reviewed lessons; render separately for QA."""
from pathlib import Path
import sys,json,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from content import PAGES

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'学习笔记/从零读懂项目_20261002';OUT.mkdir(parents=True,exist_ok=True)
RUN=ROOT/'training_runs/beginner_guide_20261002';RUN.mkdir(parents=True,exist_ok=True)
SKILL=Path('C:/Users/Amiee/.codex/plugins/cache/openai-primary-runtime/documents/26.430.10722/skills/documents')
sys.path.insert(0,str(SKILL/'scripts'))
from table_geometry import apply_table_geometry,column_widths_from_weights

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.45)
doc.core_properties.title='从零读懂量化选股项目：九项建议、实际结果与汇报准备'
doc.core_properties.subject='先举例，再介绍术语，最后核对本项目证据'
doc.core_properties.author='项目学习笔记'

def font(name,size,bold=False,color='253440'):
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size)
    st.font.bold=bold;st.font.italic=False;st.font.color.rgb=RGBColor.from_string(color)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    pf=st.paragraph_format;pf.line_spacing=1.0;pf.space_before=Pt(0);pf.space_after=Pt(4)
    pf.widow_control=True
for args in [('Normal',12),('Title',21,True,'22577A'),('Subtitle',11,False,'647582'),
    ('Heading 1',16,True,'22577A'),('Heading 2',13,True,'22577A'),('Caption',10,False,'647582'),('Quote',12,False,'22577A')]:font(*args)
for name,before,after in [('Heading 1',0,8),('Heading 2',6,3)]:
    pf=doc.styles[name].paragraph_format;pf.space_before=Pt(before);pf.space_after=Pt(after);pf.keep_with_next=True
doc.styles['Quote'].paragraph_format.left_indent=Inches(.10)
doc.styles['Quote'].paragraph_format.right_indent=Inches(0)

header=sec.header.paragraphs[0];header.text='从零读懂量化选股项目  ·  学习讲义  ·  2026-10-02'
header.runs[0].font.size=Pt(9);header.runs[0].font.color.rgb=RGBColor.from_string('647582')
border=OxmlElement('w:pBdr');bottom=OxmlElement('w:bottom')
for k,v in [('val','single'),('sz','4'),('color','D9E1E6'),('space','6')]:bottom.set(qn('w:'+k),v)
border.append(bottom);header._p.get_or_add_pPr().append(border)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('先懂例子，再会术语  ·  ');r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('647582')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def table(headers,rows,weights):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
    for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in c.paragraphs:
                pf=para.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(1);pf.line_spacing=1.0
                if i==0:pf.keep_with_next=True
                for r in para.runs:r.font.size=Pt(11);r.font.bold=i==0
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F5F7');c._tc.get_or_add_tcPr().append(sh)
    apply_table_geometry(t,column_widths_from_weights(weights or [1]*len(headers),9360),table_width_dxa=9360,indent_dxa=120,
        cell_margins_dxa={'top':60,'bottom':60,'start':120,'end':120})
    return t

def links(items):
    for label,url in items:
        para=doc.add_paragraph();para.paragraph_format.space_after=Pt(4)
        link=OxmlElement('w:hyperlink');link.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True))
        r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'22577A');pr.append(color)
        u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr)
        txt=OxmlElement('w:t');txt.text=label;r.append(txt);link.append(r);para._p.append(link)

md=['# 从零读懂量化选股项目','', '2026-10-02｜九项建议、实际结果与10月9日汇报准备','']
for i,lesson in enumerate(PAGES,1):
    if i==1:
        doc.add_heading('从零读懂量化选股项目',0)
        doc.add_paragraph('给软件工程学生的金融入门讲义｜2026-10-02','Subtitle')
    heading=doc.add_heading(f'{i:02d}  {lesson["title"]}',1)
    heading.paragraph_format.page_break_before=i>1
    # Named bookmarks make navigation possible in Word without adding a long contents page.
    mark=OxmlElement('w:bookmarkStart');mark.set(qn('w:id'),str(i));mark.set(qn('w:name'),f'lesson_{i}')
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(i));heading._p.insert(0,mark);heading._p.append(end)
    md.extend([f'## {i:02d} {lesson["title"]}',''])
    for block in lesson['blocks']:
        kind=block[0]
        if kind=='table':
            table(*block[1:]);headers,rows=block[1:3]
            md.extend(['| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|'])
            md.extend('| '+' | '.join(str(v).replace('\n','<br>') for v in row)+' |' for row in rows);md.append('')
        elif kind=='links':
            links(block[1]);md.extend(f'[{label}]({url})\n' for label,url in block[1])
        else:
            txt=block[1]
            if kind=='h':doc.add_heading(txt,2);md.extend(['### '+txt,''])
            elif kind=='say':
                para=doc.add_paragraph(style='Quote')
                if i!=26:para.add_run('汇报时可以这样说：').bold=True
                para.add_run(txt);md.extend(['> '+txt,''])
            elif kind=='check':
                para=doc.add_paragraph();para.add_run('停一下，试着解释：').bold=True;para.add_run(txt)
                md.extend(['**停一下，试着解释：**'+txt,''])
            else:doc.add_paragraph(txt);md.extend([txt,''])
filename='从零读懂量化项目_九项建议与10月9日汇报讲义.docx'
target=OUT/filename;doc.save(target)
(OUT/'从零讲义_可检索版.md').write_text('\n'.join(md),encoding='utf-8')
inputs=['training_runs/teacher9_20260930/final_summary.json','training_runs/teacher9_20260930/output_verification.json',
    'training_data/audit_20261002/inventory.json','training_data/audit_20261002/supplementary_audit.json']
manifest={'docx':str(target.relative_to(ROOT)),'lessons':len(PAGES),'characters':sum(len(str(p)) for p in PAGES),
    'docx_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),
    'inputs':{s:hashlib.sha256((ROOT/s).read_bytes()).hexdigest() for s in inputs},
    'section_titles':[p['title'] for p in PAGES],'layout_qa':'pending render and all-page visual review'}
(RUN/'guide_structure.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'file':str(target),'sections':len(PAGES),'characters':manifest['characters']},ensure_ascii=False))
