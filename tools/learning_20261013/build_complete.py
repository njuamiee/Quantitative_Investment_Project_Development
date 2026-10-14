"""Build a complete textbook and a concise project companion from audited artifacts."""
from pathlib import Path
import csv,json,re,sys,hashlib,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[2]
HERE=Path(__file__).resolve().parent
OUT=ROOT/'学习笔记/20261013_完整项目讲义'
RUN=ROOT/'training_runs/research_20261011'
AUD=ROOT/'training_runs/learning_20261013'
SKILL=Path('C:/Users/Amiee/.codex/plugins/cache/openai-primary-runtime/documents/26.1011.11113/skills/documents')
sys.path.insert(0,str(SKILL/'scripts'))
from table_geometry import apply_table_geometry,column_widths_from_weights
load=lambda p:json.loads(p.read_text('utf-8-sig'))
def csvrows(p):
    with p.open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
RESULT=load(AUD/'results_checked.json')
TABLES={}
def table(key,heads,rows,weights=None,size=9.5):TABLES[key]=(heads,rows,weights,size)
def pct(v):return f'{float(v):.2%}'
def num(v):return '—' if v is None else f'{float(v):.3f}'
def short(s):return re.sub(r'^x_c02_gtalpha','GT',re.sub(r'^x_c01_Alpha','A',s)).removeprefix('x_')


def prepare_tables():
    table('navigation',['章节','主题','读完应能解释'],[
      ('01—05','全局 数据与股票池','价格和资料怎样变成可以使用的研究样本'),
      ('06—09','因子 IC和筛选','旧40与新60/100怎样选 为什么某项被去重'),
      ('10—17','样本 标签 模型与验证','时间如何进入模型 为什么每日分数变化'),
      ('18—23','交易账本 仓位与HMM','计划怎样变成交 状态概率怎样影响决策'),
      ('24—27','评价 主结果与稳定性','哪些结论得到支持 哪些尚未通过验收'),
      ('28—32','分钟 因子择时 LLM Jev和Kronos','各分支做了什么 用了什么数据 得到什么证据'),
      ('33—38','缺口 新协议 复现 汇报与逐窗IC','下一步怎样补齐 如何解释每个季度的稳定性'),
      ('附录','全部40项 24公式 问题和任务表','可以直接查明细 不需要等下一次交付')],[1,2.1,3.4])
    table('modules_main',['模块与主要函数','输入与处理','输出与上下游'],[
      ('audit_sync.inventory/sync','远端目录与旧清单；只读核查和下载','快照、manifest、哈希；供prepare_data'),
      ('prepare_data.main','原始/后复权行情、供应商因子；复用旧prices','prices、model_panel、actions；供状态、模型、账本'),
      ('regimes.main/filtered','池内历史市场统计；月度拟合、日度前向递推','regime_probabilities；供条件评分与仓位'),
      ('train_models.screening/fit_predict','训练块与验证/测试块；筛选与模型拟合','模型对象、因子清单、各评分列'),
      ('train_models.choose/main','验证评分、价格、规则网格；模拟与选择','selection、schedule、predictions；供外层账户'),
      ('ledger.Market.run/metrics','评分、日程、价格、公司行为；逐日执行','现金、持股、成交、净值和指标'),
      ('evaluate_models.main','固定预测与账户；基准、敏感性、bootstrap','比较表、分段指标、区间估计'),
      ('verify/final_checks','合成情景与已存真实产物；断言核对','检查记录；不直接产生收益'),
      ('requirements_matrix/report_assets/build_report','来源与结果；任务映射、图表、报告','可读说明及证据入口')],[1.9,2.15,2.2],9)
    table('toy_panel',['日期','股票','因子1','因子2','未来收益','目标y'],[
      ('D1','A','0.5000','0.1667','+5%','0.5000'),('D1','B','0.1667','−0.1667','−2%','−0.1667'),('D1','C','−0.1667','0.5000','+1%','0.1667'),
      ('D2','A','−0.1667','0.5000','−1%','−0.1667'),('D2','B','0.5000','0.1667','+4%','0.5000'),('D2','C','0.1667','−0.1667','+2%','0.1667')],[.7,.65,1.1,1.1,1.05,1.2])
    table('horizons',['目录','信号','计划入场','标签终点','计划执行'],[
      ('h1','t收盘','t+1开盘','t+2开盘','次日再下一开盘退出'),
      *[(f'h{h}','t收盘','t+1开盘',f't+{h}收盘',f'第{h}个后续交易日收盘退出') for h in [2,5,10,20]]],[.55,.7,1,1,2.4])
    table('q4_windows',['环节','信号日期','标签与数量'],[
      ('初始训练','2024-02-29至2026-03-23','标签最晚2026-03-30；499日，49,881行'),
      ('验证','2026-03-31至2026-09-30','126日；跨入测试的目标置缺失'),
      ('最终拟合','2024-09-02至2026-09-22','标签最晚2026-09-30；499日，49,872行'),
      ('外层预测','2026-10-08至2026-10-09','当前快照2日、200行，目标尚未成熟')],[1,2.3,2.9])
    table('lgb_params',['参数','当前值','它限制什么'],[
      ('n_estimators / learning_rate','180 / 0.035','树的轮数和每次修正幅度'),('num_leaves / max_depth','15 / 5','分区个数和路径深度'),
      ('min_child_samples','160','每个叶子需要足够多样本'),('reg_lambda','12','叶子输出的L2正则'),('colsample_bytree','0.8','每棵树使用部分特征'),('输入维度 / 标签','最多100 / 截面收益排名','实际维度由当窗筛选名单决定')],[2.35,1.05,2.8])
    table('models',['评分列','实际机制','输入与主要设定'],[
      ('ridge60','线性Ridge','前60项；α=180'),('lgbm100','梯度提升回归树','前100项；180轮、15叶'),('ranker60','LambdaRank','前60项；同日分组、十档标签'),
      ('catboost60','提升树','前60项；180轮、深度5'),('xgboost60','提升树','前60项；180轮、深度3'),('extra60','随机化树集成','前60项；120树、深度6'),
      ('llm_ridge','LLM特征消融','24个本地提案；Ridge α=180'),('conditional_ic','状态条件因子权重','前60项及3状态概率'),('ic','训练期平均相关加权','前60项；权重绝对和归一'),('ensemble','四模型日排名平均','Ridge、LGB、CatBoost、XGBoost')],[1.4,1.65,3.1])
    scores=csvrows(AUD/'2026Q4_daily_predictions.csv')
    table('daily_scores',['日期','股票','lgbm100评分','参数版本'],[[x['date'],x['stockid'],f"{float(x['lgbm100']):.6f}",'2026Q4固定模型'] for x in scores],[1.4,1.3,1.3,2])
    matches=load(RUN/'h5/hmm_matched_exposure.json')
    table('hmm_results',['HMM方案','Sharpe差值','Calmar差值','最大回撤差值'],[[x['matched_policy'].replace('hmm',''),num(x['delta_sharpe_hmm_minus_fixed']),num(x['delta_calmar_hmm_minus_fixed']),f"{x['delta_drawdown_hmm_minus_fixed']*100:+.2f}个百分点"] for x in matches],[1.9,1.2,1.2,1.8])
    table('main_results',['期限','年化收益','Sharpe','Calmar','最大回撤'],[[f'{h}日',pct(RESULT[str(h)]['main_period_metrics']['annual_return']),num(RESULT[str(h)]['main_period_metrics']['sharpe_rf15']),num(RESULT[str(h)]['main_period_metrics']['calmar']),pct(RESULT[str(h)]['main_period_metrics']['max_drawdown'])] for h in [1,2,5,10,20]],[.8,1.3,1,1,1.3])
    table('candidate_results',['预测/持有','模型','Sharpe','Calmar','最大回撤'],[[f"{h}/{RESULT[str(h)]['candidate_metrics'].get('holding_period',h)}日",RESULT[str(h)]['candidate_metrics']['model'],num(RESULT[str(h)]['candidate_metrics']['sharpe_rf15']),num(RESULT[str(h)]['candidate_metrics']['calmar']),pct(RESULT[str(h)]['candidate_metrics']['max_drawdown'])] for h in [1,2,5,10,20]],[1.15,1.4,.95,.95,1.15])
    table('kronos_results',['期限','平均Rank IC','截面数'],[[f"{x['horizon']}日",f"{x['mean_rank_ic']:.4f}",x['dates']] for x in load(RUN/'kronos_results.json')['results']],[1.3,2.8,1.5])
    table('missing_fields',['缺少什么','影响的实验或真实性'],[
      ('历史证券主表与代码沿革','上市退市、ST、代码迁移与生存者偏差'),('精确交易状态和限价参考','替代涨跌停近似，核实停牌可成交性'),('分红到账与股份可卖日','替代除权日即时到账研究情景'),('因子发布时间和历史版本','证明每个历史时点真正可得'),('历史行业与风格资料','PIT行业市值中性化，避免今天标签回填过去'),('文本归档及带日期的人工标签','真实Jev分类、评分、校准和增益'),('宏观历史版本及融资价差','丰富HMM观测并控制发布时间'),('订单簿成交日志与真实费率','验证排队、冲击、费用与实际执行')],[2.15,4])
    old=csvrows(ROOT/'training_runs/teacher_answers_20261012/selected40_2026Q3_details.csv')
    old=sorted(old,key=lambda x:float(x['selected_order']))
    for key,part in [('old40a',old[:20]),('old40b',old[20:])]:
        table(key,['序号','因子','覆盖率','平均IC','ICIR','Ridge β'],[[int(float(x['selected_order'])),short(x['feature']),pct(x['coverage']),f"{float(x['mean_ic']):+.5f}",f"{float(x['icir']):+.4f}",f"{float(x['ridge_coefficient']):+.6f}"] for x in part],[.55,1.5,.95,1.1,1.1,1.3],9.3)
    llm=load(RUN/'llm_factor_proposals.json')
    for suffix,part in [('a',llm[:12]),('b',llm[12:])]:
        table('llm_formulas_'+suffix,['因子名','表达式'],[[x['id'].removeprefix('llm_'),x['expression']] for x in part],[2.9,3.3],9.3)
    questions=[
      ('Q01','是否2024年10月只切一次','逐季度滚动；每个季度独立训练验证和外层预测。见12章。'),
      ('Q02','train/val/test是什么日期','12章列最新h5 Q4真实日期；14章列旧40β窗口。'),
      ('Q03','叶子数学习率树数怎样定','新版固定保守参数，未证明最优；旧版验证早停。见13章。'),
      ('Q04','Ridge及正则权重','旧α=10，新ridge60为180；最小二乘加L2惩罚。见14章。'),
      ('Q05','是否考虑共线性','旧日相关0.85、新面板相关0.92去重，另有Ridge正则；不等于消除。见08—09章。'),
      ('Q06','调仓频率','每日尝试新批，h5每批t+5收盘计划退出，季度换模型。见19章。'),
      ('Q07','两模型回撤为何不同','关系、持仓路径和暴露不同；应对齐窗口逐股归因，不凭曲线猜因果。见26章。'),
      ('Q08','先选因子选股暂缓其他','新简化验收先IC、Ridge、LGB；扩展保留独立结果。见34章。'),
      ('Q09','每边1bp暂不算滑点','已做固定信号1bp敏感性；完整按1bp重选型尚未执行。见20章。'),
      ('Q10','标签和持有一致','按用户最新确认沿用t+1开盘至t+5收盘；风险/成交偏离单独记录。见11、19章。'),
      ('Q11','后复权价格','新版本地价格特征和标签已使用；真实现金账本用原始价格。见05章。'),
      ('Q12','大回撤和错误行情检查','已定位峰谷与最差日；全部逐事件外部交叉核验未完成。见26章。'),
      ('Q13','市场行业风格中性','完整PIT版本未完成；34章给出残差和持仓双层方案。'),
      ('Q14','嵌套滚动并拼OOS','现有季度内层验证与连续外层账户；老师月度新协议待执行。见12、34章。'),
      ('Q15','隔离带和标签期限','按实际label_exit剔除跨界；不能只判断信号日早于分界。见12章。'),
      ('Q16','锁定参数不回看测试','现有验证规则有明确网格；历史研究已多次查看，需新冻结期。见16、34章。'),
      ('Q17','重训必要性和频率','当前季度，短窗按月；更频繁不保证更好，480组近期未通过。见27章。'),
      ('Q18','每窗重新选因子','旧每季度初选；新初训和最终拟合分别重筛。见08—09章。'),
      ('Q19','行业数据PIT','需要生效与可得时间，不能用今天行业回填历史。见33—34章。'),
      ('Q20','每窗IC均值标准差最差轮','38章逐季列出旧LGB40、Ridge40、新h5自动评分及季度稳定性统计。'),
      ('Q21','40怎么选出来','304候选、质量门槛、绝对IC前100、日相关0.85去重、最多40；并非最优证明。见08章和两页明细。'),
      ('补充1','100股票40因子怎么训练','每天横截面按行拼，多天共同拟合；30,000×40例见10章。'),
      ('补充2','09.2右列是哪天数据','旧Ridge β，625日62,466行；日期及IC与β不同窗口见14章。'),
      ('补充3','每天重训LightGBM吗','季度重训、每天换输入；Q3/Q4已存模型推理复核见17章。')]
    table('teacher_a',['编号','老师的问题','本项目答案与位置'],questions[:11],[.6,1.75,3.8],9.3)
    table('teacher_b',['编号','老师的问题','本项目答案与位置'],questions[11:],[.6,1.75,3.8],9.1)
    qics=csvrows(ROOT/'training_runs/teacher_answers_20261012/quarterly_ic.csv')
    keyed={(x['version'],x['model'],x['quarter']):x for x in qics}
    quarters=sorted({x['quarter'] for x in qics})
    table('quarter_ic',['季度','旧LGB40','旧Ridge40','新h5自动评分'],[[q]+[f"{float(keyed[v,m,q]['mean_ic']):+.5f}" for v,m in [('old','count_40'),('old','ridge_40'),('latest_h5','score')]] for q in quarters],[1.2,1.5,1.5,1.7])
    stats=csvrows(ROOT/'training_runs/teacher_answers_20261012/quarterly_ic_stability.csv')
    stats={(x['version'],x['model']):x for x in stats}
    table('quarter_stability',['模型','季度均值的均值','季度均值标准差','最差季度'],[[label,f"{float(stats[k]['mean_of_fold_ic']):.5f}",f"{float(stats[k]['std_of_fold_ic']):.5f}",stats[k]['worst_fold']] for k,label in [(('old','count_40'),'旧LGB40'),(('old','ridge_40'),'旧Ridge40'),(('latest_h5','score'),'新h5自动')]], [1.3,1.65,1.65,1.2])
    tasks=csvrows(RUN/'requirements_matrix.csv')
    supplements={
      'R02':'每日更新池和因子值；模型和入选身份季度更新，HMM参数月更。04、09、17章。',
      'R04':'筛选已实现；新版近似相关与旧严格Rank IC差异必须注明。07—09章。',
      'R12':'主5bp与固定信号敏感性已做；老师1bp重选型待做。20章。',
      'R14':'全期探索部分通过，自动选择及近期稳定性失败。25、27章。',
      'R24':'24项本地提案与受限DSL已实现；没有持续自主挖掘服务。30章。',
      'R28':'本次一次交付完整教材、导读及全部证据附件。',
      'R29':'材料已备，会议预约与发送需本人安排。',
      'R30':'完整教材和手算题已交付；个人掌握需自行推演。'}
    chapters=['04','04 09 17','06','07 08 09','12 17 18','13 14 15','12 18','11','24 25','24','26','20','21 26','25 27','22','22','23','23 29','28','28','28','05 20','33','30','31','32','29','全书','37','37']
    rows=[]
    for i,x in enumerate(tasks):
        desc=supplements.get(x['id'],x['status']+'。见'+chapters[i]+'章。')
        rows.append([x['id'],x['requirement'],desc])
    for key,part in [('tasks_a',rows[:10]),('tasks_b',rows[10:20]),('tasks_c',rows[20:])]:table(key,['任务','要求','实际状态与解决位置'],part,[.55,2.05,3.55],9.3)
    table('modules_extra',['模块','输入与功能','主要输出'],[
      ('short_windows','面板与两HMM版本；月度短窗','short_windows/all_results、folds、summary'),('risk_refinement','既有评分；风险缩放与保护网格','h*/risk_refinement、candidate_*'),
      ('portfolio_extension','h10/h20评分；预测持有分离','portfolio_extension；旧候选留aligned_holding'),('intraday','分钟记录与前日池；IC和聚合特征','minute_ic、minute_factor_*'),
      ('legacy_extensions/refresh_minute_accounts','旧沪深/分钟评分；重记账','shsz_cost5、minute_execution_*'),('factor_timing','面板、已实现因子收益与概率','state_factor_matrix、factor_timing_*'),
      ('factor_dsl','受限表达式与历史列；求值','24项本地输入进入面板'),('jev_adapter','带版本时间的证据；构造与校验','jev_offline_request、jev_status'),
      ('kronos_probe','公开权重与128日序列；预测','kronos_predictions、kronos_results'),('period_audit/recent_diagnosis','已有净值和评分；分段诊断','all_strategy_period_audit、recent_signal_diagnosis'),
      ('supplement_audit/inspect_supplement','补充目录与文件内容；检查','补充清单、缺口和源文件信息')],[1.85,2.15,2.2],9)


class Writer:
    def __init__(self,name):
        self.name=name;self.d=Document();self.md=[];self.citations=[]
        # Bundled document defaults can carry a decorative Title border.
        for border in list(self.d.styles.element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
        sec=self.d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
        sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
        sec.header_distance=sec.footer_distance=Inches(.4)
        for key,size,bold in [('Normal',11.5,False),('Title',22,True),('Subtitle',11,False),('Heading 1',16,True),('Heading 2',12.5,True),('Caption',9,False)]:
            st=self.d.styles[key];st.font.name='Arial';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0)
            st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
            st.paragraph_format.line_spacing=1.08;st.paragraph_format.space_after=Pt(6)
            if key.startswith('Heading'):st.paragraph_format.space_before=Pt(9);st.paragraph_format.keep_with_next=True
        h=sec.header.paragraphs[0];h.text='多期限量化选股项目  原理  代码  证据';h.style='Caption'
        f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('2026年10月13日  |  ').font.size=Pt(9)
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
        self.d.core_properties.title=name;self.d.core_properties.author='项目科研学习记录'
    def paragraph(self,text,style=None):
        p=self.d.add_paragraph(text,style);self.md.append(text);return p
    def grid(self,key):
        heads,rows,weights,size=TABLES[key];t=self.d.add_table(rows=1,cols=len(heads));t.style='Table Grid'
        for c,s in zip(t.rows[0].cells,heads):c.text=str(s)
        for row in rows:
            for c,s in zip(t.add_row().cells,row):c.text=str(s)
        compact=key in ['models','main_results','candidate_results']
        apply_table_geometry(t,column_widths_from_weights(weights or [1]*len(heads),9360),table_width_dxa=9360,indent_dxa=120,cell_margins_dxa={'top':45 if compact else 65,'bottom':45 if compact else 65,'start':90,'end':90})
        pr=t._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        old=pr.find(qn('w:tblBorders'))
        if old is not None:pr.remove(old)
        pr.append(borders)
        for i,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for j,c in enumerate(row.cells):
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in c.paragraphs:
                    p.paragraph_format.line_spacing=1.03;p.paragraph_format.space_after=Pt(.5 if compact else 2);p.paragraph_format.space_before=Pt(.5 if compact else 2)
                    s=c.text
                    if len(s)<=12 and key not in ['llm_formulas_a','llm_formulas_b','modules_main','modules_extra']:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:r.font.size=Pt(size);r.font.bold=i==0;r.font.color.rgb=RGBColor(0,0,0)
                if i==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EFF5');c._tc.get_or_add_tcPr().append(sh)
        clean=lambda s:str(s).replace('|','\\|').replace('\n',' / ')
        self.md.append('\n'.join(['| '+' | '.join(map(clean,heads))+' |','| '+' | '.join(['---']*len(heads))+' |']+['| '+' | '.join(map(clean,row))+' |' for row in rows]))
    def source(self,line):
        file,needle,meaning=line.split('|',2);lines=(ROOT/file).read_text('utf-8-sig').splitlines()
        hits=[i for i,s in enumerate(lines,1) if needle in s]
        if not hits:raise ValueError('Source needle missing: '+line)
        text=f'源码定位 {file}:{hits[0]}。{meaning}'
        if self.d.paragraphs and self.d.paragraphs[-1].style.name=='Caption' and self.d.paragraphs[-1].text.startswith('源码定位 '):
            self.d.paragraphs[-1].add_run('\n'+text);self.md.append(text)
        else:self.paragraph(text,'Caption')
        self.citations.append({'file':file,'line':hits[0],'needle':needle})
    def render(self,content):
        lines=content.splitlines();i=0
        while i<len(lines):
            line=lines[i].strip();i+=1
            if not line:continue
            if line.startswith('```'):
                lang=line[3:];block=[]
                while i<len(lines) and not lines[i].strip().startswith('```'):block.append(lines[i]);i+=1
                i+=1;p=self.d.add_paragraph();p.paragraph_format.line_spacing=1.0
                for j,s in enumerate(block):
                    r=p.add_run(('\n' if j else '')+s);r.font.name='Consolas';r.font.size=Pt(9)
                self.md.append('```'+lang+'\n'+'\n'.join(block)+'\n```');continue
            if line.startswith('# '):self.paragraph(line[2:],'Title');continue
            if line.startswith('## '):
                p=self.d.add_heading(line[3:],1);p.paragraph_format.page_break_before=True;self.md.append(line);continue
            if line.startswith('### '):self.d.add_heading(line[4:],2);self.md.append(line);continue
            if line.startswith('@source '):self.source(line[8:]);continue
            if line.startswith('@figure '):
                name=line[8:];self.d.add_picture(str(OUT/'配图'/name),width=Inches(6.45));self.md.append('!['+name+'](配图/'+name+')');continue
            if line.startswith('@'):self.grid(line[1:]);continue
            link=re.fullmatch(r'\[(.*?)\]\((https?://.*?)\)',line)
            if link:
                p=self.d.add_paragraph();hy=OxmlElement('w:hyperlink');hy.set(qn('r:id'),p.part.relate_to(link[2],RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=link[1];r.append(t);hy.append(r);p._p.append(hy);self.md.append(line);continue
            self.paragraph(line)
        for border in list(self.d.element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
        self.d.save(OUT/(self.name+'.docx'));(OUT/(self.name+'.md')).write_text('\n\n'.join(self.md),'utf-8')


COMPANION='''# 从零读懂量化项目完整导读

项目全貌 核心方法 研究结论和文件导航

本导读用于先看清整个项目；配套《从零讲义 全项目完整详解》一次提供所有模块的原理、数字案例、公式、源码位置，以及40项因子明细和老师全部问题。两份材料均已完整交付，不按课次等待后续内容。

项目每天对动态股票池评分，预测未来1、2、5、10、20交易日期限的相对收益，再用现金股数账本评价。数据快照截至2026年10月9日，解释按10月13日当前源码核对。

最关键结论是：多天横截面按行拼接训练，主模型每季度拟合、每日预测；旧表最右列是多日联合拟合的Ridge系数；模型评分、交易执行和账户收益是不同层。当前自动验证选择尚未同时达到Sharpe大于1、Calmar大于2。

### 完整教材导航

@navigation

## 01 数据流与依赖

@figure 全局数据流.png

prepare_data产生每日股票池、特征与未来标签；regimes产生过去信息驱动的状态概率；train_models做各期限筛选、拟合和验证选择；ledger模拟执行；evaluate和分段审计形成结果。验证选择也调用同一个ledger。

输入快照在training_data，结果在training_runs，当前主代码在tools/research_20261011。新流程仍依赖9月30日与10月8日的历史数据和部分函数。

## 02 每个主模块负责什么

@modules_main

先按run_local读调用顺序，再跟随model_panel到screening、fit_predict、choose，最后读Market.run。完整教材第03章和附录给出全部扩展模块及结果文件。

## 03 样本因子模型和时间轴

100股票×40因子×300天且无缺失，训练X为30,000行×40列；每个预测日输入100行，输出100个分数。date和stockid用于组织记录，不作为当前LightGBM输入。时间信息来自历史滚动特征、切分日期和标签成熟边界。

旧40因子版本：304候选经覆盖、IC日期和绝对IC筛选进入前100，再按平均日相关0.85去重，最多40。当前多期限版本：335候选，覆盖70%，近似相关排序、面板相关0.92去重；LGB最多100，多数模型60。两版定义必须分别说。

旧09.2的Ridge β使用2023-11-21至2026-06-23的62,466行最终拟合，标签最晚6月30日。IC等统计用更早初始训练段。完整教材第08、14章和两页40项明细逐条解释。

@q4_windows

训练窗、验证窗、重拟合窗、预测期限和持有期限是不同时间长度。当前主模型按季度更新，HMM参数按月，股票池和评分每日更新；短窗口实验另按月更新。

## 04 模型怎样变成实际持仓

当前训练10类评分：线性Ridge、LightGBM、Ranker、CatBoost、XGBoost、ExtraTrees、LLM专用Ridge、IC加权、条件IC和集成。验证期先按Rank IC选前三，再比较90个账户组合；目标是min(Sharpe,Calmar/2)。

模型只给排序分数。账本按t收盘信号、t+1开盘入场、h5在t+5收盘计划退出；每天开新批次，每批预算不超过净值1/5。实际受现金、10%单股上限、T+1、最低100股、过去ADV20的1%共同买卖容量及不可成交条件约束。

HMM提供状态概率，可用于软仓位、条件因子评分和独立因子配置。它必须与固定仓位、纯波动和趋势规则比较。状态识别合理也可能映射成无效交易动作，不能预设必然有正增益。

主账户使用5bp每边；老师1bp每边简化协议尚待完整重选型。当前已有固定评分1bp敏感性。后复权用于标签和本地价量特征，原始价格用于现金账本；除权日即时到账属于研究假设。

## 05 当前结果和证据强度

@main_results

自动选择统一区间为2024-10-08至2026-09-30，共485日，5bp每边、Sharpe无风险年率1.5%。没有一个期限同时通过两道门槛。

全历史探索中有达标候选，但近期失效。480组短窗月度实验：全期41组通过，近一年0组、近半年0组、三段同时通过0组。历史挑出的最好结果不能覆盖自动选择失败。

分钟因子完成少量日期的研究与执行压力；因子择时尚未证明稳定增益；24项LLM提案已进入本地研究；Jev仅离线；Kronos仅9个稀疏截面推理。完整教材28—32章给出每项设计、公式、数值和限制。

6项模型推理检查、已有金融约束测试支持特定程序性质；它们不保证行情全部正确，也不保证未来收益。

## 06 尚未完成事项和老师问题入口

已实现：数据与每日股票池、五期限监督模型、月度HMM、连续账户、费用和基准对照，以及各独立研究分支。尚未完成：老师1bp简化新协议、完整PIT行业风格中性化、全部异常行情外部核验、真实Jev文本评估、实际到账与成交校验、生产调度及前瞻收益验收。

最新明确事项：你已接受t+5或t+6内部一致，因此沿用当前t+5。下一轮可以先固定简化IC、Ridge和LGB基线，再做月度嵌套验证、中性化和独立HMM增益比较。完整方案在教材34章。

老师10月12日21个问题、后续3个时间轴问题、10月9日与项目规划30项任务，都已在完整教材附录逐项回答并标注状态。Q20相关逐季IC表、40项去重过程及7329条已识别文件缺口在核对资料中。

已经发现screening相关口径、recent_diagnosis字段名、验证失败仍选候选和最终重筛名单变化等问题。教材解释实际代码，不通过修改措辞把未完成事项写成已验证。

## 07 如何阅读源码和复现

完整教材35章给出本地复现命令和检查顺序。研究环境是.venv_research_20261011/Scripts/python.exe；先跑verify.py，再根据需要运行run_local.ps1。默认本地，-Sync才联网盘点下载，Jev默认离线。

各章节的源码定位包含真实文件、函数和行号，行号对应本次代码哈希。核对资料还保存335项候选目录、全部24项公式、旧40项详细统计、训练窗口、模型真实日评分、任务表、文件缺口与字段缺口。

汇报时依次解释目标、训练矩阵、标签与日期、因子筛选、模型、交易账本、实验结论、局限和下一步。最重要的能力是沿一条记录追踪其输入何时可得、标签何时成熟、分数如何决定交易和净值。

主代码、旧快照与部分历史依赖都需要保留。远程仓库不包含所有原始数据与本机环境，单独克隆摘要并不能直接重现训练。改变数据或参数应建立新研究版本，保留原输入和结果哈希。
'''


def main():
    (OUT/'配图').mkdir(parents=True,exist_ok=True)
    for image in (AUD/'figures').glob('*.png'):shutil.copy2(image,OUT/'配图'/image.name)
    prepare_tables()
    raw=(HERE/'full_handbook.md').read_text('utf-8')
    raw=raw.replace('@llm_formulas','@llm_formulas_a\n\n## 附录三续 本地LLM因子公式第13至24项\n\n@llm_formulas_b')
    writers=[Writer('01_从零读懂量化项目_完整导读'),Writer('02_从零讲义_全项目完整详解')]
    writers[0].render(COMPANION);writers[1].render(raw)
    support=OUT/'核对资料';support.mkdir(exist_ok=True)
    copies={
      AUD:['source_index.json','model_examples.json','old_table_answer.json','inference_checks.json','results_checked.json','2026Q3_daily_predictions.csv','2026Q4_daily_predictions.csv'],
      RUN:['factor_catalog.csv','llm_factor_proposals.json','missing_files.csv','missing_fields_and_documents.csv','requirements_matrix.csv','kronos_results.json'],
      ROOT/'training_runs/teacher_answers_20261012':['selected40_2026Q3_details.csv','selected40_all_quarters.csv','selection_trace_2026Q3.csv','selection_summary.json','quarterly_ic.csv','quarterly_ic_stability.csv','drawdown_diagnostics.csv']}
    for folder,names in copies.items():
        for name in names:shutil.copy2(folder/name,support/name)
    shutil.copy2(RUN/'sources.md',support/'研究来源.md')
    old_index=ROOT/'学习笔记/20261013_从零学习当前项目/核对资料/代码索引.md'
    if not old_index.exists():
        old_index=ROOT/'历史学习资料/20261013_分课初稿/20261013_从零学习当前项目/核对资料/代码索引.md'
    if old_index.exists():shutil.copy2(old_index,support/'代码索引.md')
    (OUT/'章节索引.md').write_text('# 完整讲义章节\n\n'+'\n'.join('- '+x[3:] for x in raw.splitlines() if x.startswith('## ')),'utf-8')
    (support/'本文源码引用.json').write_text(json.dumps(writers[1].citations,ensure_ascii=False,indent=2),'utf-8')
    manifest={p.relative_to(OUT).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in OUT.rglob('*') if p.is_file() and p.name!='delivery_manifest.json'}
    (OUT/'delivery_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),'utf-8')
    print(json.dumps({'output':str(OUT),'chapters':raw.count('\n## '),'source_citations':len(writers[1].citations),'files':len(manifest)},ensure_ascii=False))

if __name__=='__main__':main()
