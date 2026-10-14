"""Two modular teaching documents, built from the read-only project audit."""
from pathlib import Path
import json,csv,sys,hashlib,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[2]
AUD=ROOT/'training_runs/learning_20261013'
OUT=ROOT/'学习笔记/20261013_从零学习当前项目'
ASSET=OUT/'配图';ASSET.mkdir(parents=True,exist_ok=True)
SKILL=Path('C:/Users/Amiee/.codex/plugins/cache/openai-primary-runtime/documents/26.430.10722/skills/documents')
sys.path.insert(0,str(SKILL/'scripts'))
from table_geometry import apply_table_geometry,column_widths_from_weights
load=lambda n:json.loads((AUD/n).read_text('utf-8'))
INDEX=load('source_index.json');EX=load('model_examples.json');OLD=load('old_table_answer.json');RESULT=load('results_checked.json')
BLUE='#173F55';TEAL='#2C756D';GRAY='#586974'

def source(file,needle):
    path=ROOT/file
    lines=path.read_text('utf-8-sig').splitlines()
    matches=[i for i,x in enumerate(lines,1) if needle in x]
    if not matches:raise ValueError((file,needle))
    return file+':'+str(matches[0])




class Guide:
    def __init__(self,stem,header):
        self.stem=stem;self.doc=Document();self.md=[]
        sec=self.doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
        sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
        sec.header_distance=sec.footer_distance=Inches(.4)
        for name,size,bold,color in [('Normal',12,False,'263238'),('Title',23,True,'173F55'),('Subtitle',12,False,'586974'),('Heading 1',16,True,'173F55'),('Heading 2',13,True,'2C756D'),('Caption',9,False,'586974')]:
            st=self.doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor.from_string(color)
            st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
            st.paragraph_format.line_spacing=1.08;st.paragraph_format.space_after=Pt(6)
            if name.startswith('Heading'):st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(10)
        h=sec.header.paragraphs[0];h.text=header;h.runs[0].font.size=Pt(9)
        f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('2026-10-13  |  ').font.size=Pt(9)
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
        self.doc.core_properties.title=stem;self.doc.core_properties.author='项目科研学习记录'
    def p(self,text,style=None):
        self.md.append(text);return self.doc.add_paragraph(text,style)
    def title(self,text):self.p(text,'Title')
    def h(self,text):self.doc.add_heading(text,2);self.md.append('### '+text)
    def page(self,text):
        p=self.doc.add_heading(text,1);p.paragraph_format.page_break_before=True;self.md.append('## '+text)
    def table(self,headers,rows,weights=None,size=10):
        t=self.doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
        for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
        for r in rows:
            for c,v in zip(t.add_row().cells,r):c.text=str(v)
        apply_table_geometry(t,column_widths_from_weights(weights or [1]*len(headers),9360),table_width_dxa=9360,indent_dxa=120,cell_margins_dxa={'top':65,'bottom':65})
        for i,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            if i==0:pr.append(OxmlElement('w:tblHeader'))
            for c in row.cells:
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.06
                    for r in p.runs:r.font.size=Pt(size);r.font.bold=i==0
                if i==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEF3F5');c._tc.get_or_add_tcPr().append(sh)
        cell=lambda s:str(s).replace('|','\\|').replace('\n',' / ')
        self.md.append('\n'.join(['| '+' | '.join(map(cell,headers))+' |','| '+' | '.join(['---']*len(headers))+' |']+['| '+' | '.join(map(cell,r))+' |' for r in rows]))
    def fig(self,name,caption):
        self.doc.add_picture(str(ASSET/name),width=Inches(6.45));self.p(caption,'Caption');self.md.append(f'![{caption}](配图/{name})')
    def code(self,text):
        p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(6)
        for i,line in enumerate(text.splitlines()):
            if i:p.add_run('\n')
            r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(9.5)
        self.md.append('```python\n'+text+'\n```')
    def cite(self,file,needle,meaning):self.p('源码定位：'+source(file,needle)+'；'+meaning,'Caption')
    def link(self,label,url):
        p=self.doc.add_paragraph();hy=OxmlElement('w:hyperlink');hy.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
        r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=label;r.append(t);hy.append(r);p._p.append(hy);self.md.append(f'[{label}]({url})')
    def save(self):
        self.doc.save(OUT/(self.stem+'.docx'));(OUT/(self.stem+'.md')).write_text('\n\n'.join(self.md),'utf-8')

def global_guide():
    g=Guide('01_从零读懂量化项目_全局导读','当前项目全局导读  |  先建立地图，再逐课学习')
    g.title('从零读懂量化项目\n先知道每段代码在做什么')
    g.p('面向软件工程研究生｜当前代码快照：2026-10-13｜行情截止：2026-10-09','Subtitle')
    g.h('这个项目要解决的问题')
    g.p('在每个交易日收盘后，利用当时可用的价量和因子信息，对一个动态股票池里的股票排序，研究未来1、2、5、10、20个交易日期限上的相对表现。然后把评分送入带现金、持股、交易成本和成交约束的模拟账户，检验排序能否转化为有风险约束的净收益。')
    g.h('用软件工程语言理解')
    g.p('这是一个有时间约束的数据流水线：采集与校验 → 特征表 → 监督学习 → 批量推理 → 决策规则 → 交易模拟器 → 指标与审计。训练脚本像离线批处理；模型文件像版本化构建产物；输入哈希和运行记录类似可复现构建日志。金融上的额外困难，是未来标签要等待、行情会变、交易会受限制。')
    g.h('一个小例子')
    g.p('假设今天池里有100只股票，模型给100个分数，规则拟买排名前10。即使排名第一，账户也可能因现金不足、最低股数或成交容量而买不到。股票评分、计划持仓、实际成交、账户收益是四层不同的数据，必须逐层核对。')
    g.h('先明确当前结论')
    g.p('项目已经能够进行本地滚动实验，但还不是经过前瞻验证、可直接实盘运行的稳定盈利系统。按验证规则自动选型的方案未同时达到Sharpe>1、Calmar>2；部分历史探索候选达到全期数字，近期窗口仍失败。')
    g.p('学习入口只有两份：本导读负责全局地图；《02_从零讲义_第1课》只解释样本、时间轴、训练和每日预测。其余模块按课程顺序逐课展开，不要求一次学完。','Caption')

    g.page('1  完整数据流：文件怎样变成账户结果')
    g.fig('全局数据流.png','实线表示当前数据或结果依赖；扩展分支分别研究，不能把它们的最好结果混成同一条策略。')
    g.table(['层次','核心产物与含义'],[
        ('输入','training_data/research_20261011：原始/后复权日线、供应商因子、公司行为；同时依赖9月30日和10月8日历史快照。'),
        ('学习','model_panel.parquet：每行一只股票在一个日期的特征和离线标签；models.joblib：训练后固定的模型。'),
        ('决策','h*/predictions.parquet：每天每只股票的评分；schedule.csv：由验证选择的模型配套交易设置。'),
        ('执行与证据','selected_daily/orders/trades/holdings：现金、股数、成交和净值；metrics、分段结果、核验文件供复查。')],[1,5],9.5)
    g.p('箭头顺序不是随意调用：run_local.ps1规定主流程。Jev默认禁止付费联网；Kronos推理独立运行；目前没有每日收盘自动调度与券商下单系统。','Caption')

    g.page('2  模块依赖：输入、动作、输出')
    g.table(['模块 / 函数','输入 → 动作 → 输出','上下游'],[
        ('audit_sync.inventory / sync','远端目录与旧快照 → 只读盘点、下载、哈希 → manifest与文件','准备数据前；默认复现可不联网'),
        ('prepare_data.main','行情与因子 → 股票池、历史特征、未来标签、同日排名 → prices / model_panel','依赖旧prepare.prices；供训练与状态识别'),
        ('regimes.main / filtered','市场历史统计 → 月度HMM与每日前向概率 → regime_probabilities','供条件评分、仓位上限'),
        ('train_models.screening','当前训练窗 → IC近似统计、覆盖与去重 → 因子列表','只用当前可训练区间；最多100项'),
        ('fit_predict / choose / main','训练、验证、测试块 → 拟合、验证选型、季度重拟合 → 模型/评分/设置','choose调用ledger在验证期模拟；最终评分也交给ledger'),
        ('ledger.Market.run / metrics','评分、价格、公司行为、风险设置 → 逐日现金股数记账 → 账户与指标','各实验共同调用；主链默认5bp/边'),
        ('evaluate_models.main','固定评分与账户 → 对照、费用敏感性、分段、bootstrap → 评估结果','不会把事后最好方案回写成自动选择成绩'),
        ('verify / final_checks','合成数据及真实产物 → T+1、资金、容量、边界检查 → 检查记录','检查正确性；不证明未来盈利'),
        ('requirements_matrix / report_assets / build_report','结果与任务 → 来源映射、图表、报告','展示层；不产生新的投资收益')],[1.75,2.85,1.4],9)
    g.p('独立扩展：short_windows（月度短窗）、risk_refinement与portfolio_extension（历史探索）、intraday与legacy_extensions（分钟/沪深）、factor_timing（因子收益择时）、factor_dsl（24个公式）、jev_adapter（离线文本接口）、kronos_probe（本地公开模型）。逐文件函数与行号见“核对资料/代码索引.md”。','Caption')

    g.page('3  实现状态与证据：当前能说到哪一步')
    g.table(['状态','当前证据'],[
        ('已经实现并留下产物','335项候选输入、动态Top100、5个期限、季度监督模型、月度HMM、连续账户、分钟与短窗扩展。'),
        ('已核对但范围有限','9项合成检查、5个期限产物检查；本次重载Q3/Q4模型，6项预测一致性检查通过。'),
        ('尚未完成','按老师新简化协议的完整重跑、完整PIT行业/风格中性化、全部异常价格外部核验、真实Jev文本评估、每日生产调度。'),
        ('未经验证的条件','供应商因子的历史发布版本、历史证券资格、分红真实到账和送转可卖日；旧回测不能自动保证这些条件。'),
        ('待检验的研究假说','更短训练窗一定更好、HMM一定有增益、更多因子一定更有效、正IC足以保证多头赚钱，均不能当作事实。')],[1.55,4.45],9.5)
    g.h('先看自动选择成绩，再看探索成绩')
    g.table(['预测期限','Sharpe¹','Calmar','最大回撤'],[[str(h)+'日',f"{RESULT[str(h)]['main_period_metrics']['sharpe_rf15']:.3f}",f"{RESULT[str(h)]['main_period_metrics']['calmar']:.3f}",f"{RESULT[str(h)]['main_period_metrics']['max_drawdown']:.2%}"] for h in [1,2,5,10,20]],[1.4,1.4,1.4,1.8])
    g.p('¹ 固定比较口径：2024-10-08至2026-09-30，485个交易日，5bp/边，Sharpe无风险年率1.5%；数据文件h*/main_period_metrics.json。完整快照还包含10月8/9日，两种截止日期的数字不能混用。','Caption')
    g.p('短窗月度扩展共480组：全期41组达标，近一年0组、近半年0组，三个区间同时达标0组。1bp费用敏感性已有结果，但与按1bp重新进行完整验证选择是两件不同的实验。老师接受t+5/t+6内部一致，本轮教学沿用现有t+5定义。')

    g.page('4  学习地图：每次只讲一个模块')
    g.table(['课程','需要理解的知识','学完能独立回答'],[
        ('第1课：样本与时间轴','面板数据、横截面、特征/标签、fit/predict、重训与持有周期','100只股票×40因子怎样进模型？旧表β来自何时？'),
        ('第2课：数据与股票池','成交量/额、价格复权、缺失、历史可得性与生存者偏差','为什么池子每天变？什么时候算偷看未来？'),
        ('第3课：因子与筛选','动量/反转、Rank IC、ICIR、相关去重、多重比较','40/60/100项怎么选？负IC怎样理解？'),
        ('第4课：模型与训练','Ridge正则、树与提升、排序损失、验证选择、泄漏','学习的目标是什么？参数为什么这样设？'),
        ('第5课：账户与回测','现金/股数、T+1、费用、分批持有、容量与延期成交','为什么好评分没有变成好净值？'),
        ('第6课：指标与研究结论','年化、Sharpe、Calmar、最大回撤、分段、置信区间','什么算有效证据？如何避免挑好结果？'),
        ('第7课：HMM与扩展','状态概率、风险预算、对照与消融、LLM/Jev/Kronos边界','增加模块是否真的产生样本外增益？')],[1.45,2.35,2.2],9.5)
    g.h('固定教学方式')
    g.p('每课按“直观问题 → 小数字手算 → 必要公式 → 本项目文件/函数/行号 → 自测”展开。先由你复述核心结论，再进入下一课。遇到老师追问，先给一句准确结论，再给一组数据和一处源码，不靠堆模型名回答。')
    g.h('阅读约定')
    g.p('金融知识在这里用于理解收益和交易约束；统计知识用于衡量信号、噪声和研究偏差；机器学习知识用于把历史映射变为可复核的预测。代码写出了某功能、检查通过了某性质、实验支持了某结论，三者分别陈述。')
    g.p('本次核对当前研究目录26个Python脚本、配置与流程文件及旧版必要依赖，共35个源文件；重载已存模型和读取已存结果，不重新调参、不启动SFTP同步、不调用Jev。行号以本次代码哈希为准。','Caption')
    g.save()

def first_lesson():
    g=Guide('02_从零讲义_第1课_样本与时间轴','第1课  |  样本、时间轴、模型训练与每日预测')
    g.title('从零讲义 · 第1课\n100只股票、40个因子与时间轴')
    g.p('本课只解决一个模块：历史记录如何进入模型，模型如何生成每天的新分数。','Subtitle')
    g.table(['你的问题','直接回答'],[
        ('是每天的横截面回归吗？','每天形成一个横截面；训练时把训练窗里多天的横截面按行拼接，联合拟合一套模型。每日单独拟合100条样本并非当前实现。'),
        ('时间轴怎么处理？','靠截至当日的历史特征、按日期切分、剔除未成熟标签、按季度更新。date和stockid不进入当前LightGBM特征矩阵。'),
        ('每天收盘重训LightGBM吗？','主流程每季度重训；每天更新股票池、因子值和评分。短窗口扩展另用月度训练。'),
        ('第11页09.2最右列是哪天的？','那列是旧版2026Q3的Ridge β。用2023-11-21至2026-06-23的62,466条股票日记录联合拟合，标签最晚到6月30日。它是整段历史估计出的系数。'),
        ('t+5/t+6如何处理？','接受内部一致要求。本课沿用t+1开盘→t+5收盘；明确区分5日持有与每日新建批次。')],[1.65,4.35],10)
    g.fig('训练与预测.png','旧版讲40列；最新主LightGBM实际为100列。训练与预测的组织方式相同。')
    g.p('先记住一句话：多天联合训练一套模型，当天100行输入产生100个分数；新分数不意味着新模型。')

    g.page('1  先读懂你指出的表：一行是一项因子')
    g.h('直观解释')
    g.p('旧解答09.2是一张“因子统计与模型参数表”。它的20行列出第21—40个因子；训练矩阵则是一行一只股票在某天的观察。两张表的行含义不同。最右列Ridge β表示线性模型给该因子的联合系数，LightGBM的树模型没有这组固定线性β。')
    g.table(['旧表列','含义与使用的数据窗口'],[
        ('覆盖率、平均IC、ICIR','筛因子训练窗：2023-11-21至2025-12-10，499个信号日；标签最晚2025-12-17。IC是每日截面相关再对日期平均。'),
        ('最右列Ridge β','测试前最终重拟合：2023-11-21至2026-06-23，625个有效信号日、62,466条样本；标签最晚2026-06-30。'),
        ('使用时段','2026Q3测试信号日为7月1日至9月29日；这套β在该季度固定，下一季度按新窗口重新估计。')],[1.3,4.7],10)
    g.h('数字例子与必要公式')
    beta=OLD['coefficients_21_40']['x_c02_gtalpha112']
    g.p(f'表中第21项GT112的β≈{beta:.6f}。在模型计算中，若其输入排名值从0.1变成0.2，其他输入数值保持不变，则这一项对预测分数的贡献增加约{beta*.1:.6f}。这里的分数单位来自排名标签，不能解释成股票上涨了该百分比。')
    g.p('Ridge分数 = b + Σ βⱼxⱼ；本轮b≈0.004833，α=10。拟合目标是让所有有效股票日样本的平方误差之和，再加10×Σβⱼ²，尽量小。40个系数共同估计，不能只看单因子IC就确定β。')
    g.cite('tools/research_20260930/train.py','ridge=Ridge(alpha=10)','rolling()中的联合拟合；下一行预测，随后保存ridge.npz。')
    g.p('直接证据：training_runs/teacher9_20260930/bj/2026Q3/ridge.npz、fold.json；本次从原面板重新核对样本日期与条数。模型内部的重拟合时点为“测试前”，不等于代码实际运行的系统日期。','Caption')

    g.page('2  训练矩阵：把多天横截面往下拼')
    g.h('直观解释')
    g.p('把它当成数据库表，主键为(date, stockid)。选定40个因子后，取40列作为X，未来收益排名这一列作为y。日期增加意味着增加行；股票池成员可以变。模型学习所有训练行共享的一套映射。')
    g.table(['日期','股票','因子1排名值','因子2排名值','未来收益','目标y'],[
        ('D1','A','0.5000','0.1667','+5%','0.5000'),('D1','B','0.1667','−0.1667','−2%','−0.1667'),('D1','C','−0.1667','0.5000','+1%','0.1667'),
        ('D2','A','−0.1667','0.5000','−1%','−0.1667'),('D2','B','0.5000','0.1667','+4%','0.5000'),('D2','C','0.1667','−0.1667','+2%','0.1667')],[.55,.5,1.3,1.3,1.2,1.15],9.5)
    g.p('这是教学例子：2天×3只股票×2因子，X形状为6×2，y长度为6。未来收益和y仅在历史训练时已成熟后使用；预测当天模型只接收两个因子列。','Caption')
    g.h('公式与维度')
    g.p('X = [X_D1；X_D2；…；X_DT]，M = Σₜnₜ，X∈ℝ^(M×K)。若每一天恰好100只、500天都有效，M=50,000，K=40；现实有缺失和标签剔除，条数会减少。旧表对应最终X是62,466×40；新2026Q4的LightGBM最终X是49,872×100。')
    g.p('代码按行的回归误差拟合，没有显式“每日等权”训练权重；某天有效行更多，就贡献更多训练行。同日股票与重叠收益标签有关联，50,000行不代表50,000个独立金融样本。每日IC的日期平均又是另一层统计。')
    g.cite('tools/research_20261011/train_models.py','x=train[cols].fillna(0).to_numpy(dtype=np.float32)','fit_predict()生成二维数组；普通回归分支执行m.fit(x,y)。')
    g.cite('tools/research_20261011/prepare_data.py',"ranked=panel[cols]",'main()按date做截面排名；股票池由旧prepare.prices()构建。')

    g.page('3  标签：历史答案怎样与当天输入配对')
    g.h('直观解释与手算')
    g.p('t日收盘时，我们知道当天特征，但不知道随后持有的收益。历史训练时先等该收益实现，再把它贴到t日这行后面。假设次日开盘10元、t+5收盘10.50元，则这一行的5日期限收益标签为5%。')
    g.p('在上一页D1的3只股票中，未来收益为[5%, −2%, 1%]，从小到大名次为[3,1,2]。代码用百分位名次减0.5，因此目标是[0.5, −0.1667, 0.1667]。同日原始因子也做类似排名。它们是相对位置，不是上涨概率。')
    g.h('必要公式')
    g.p('rᵢ,ₜ⁽⁵⁾ = Cᵢ,ₜ₊₅(adj) / Oᵢ,ₜ₊₁(adj) − 1\nyᵢ,ₜ = 当日有效收益的百分位排名 − 0.5\nxᵢ,ₜ,ⱼ = 当日因子j有效值的百分位排名 − 0.5')
    g.p('adj表示后复权价格。3只股票例子中，中间名次是0.1667；减0.5并不保证样本均值严格为0。有并列时使用平均名次，缺值有各自的有效样本集合；模型阶段填0是中性近似。')
    g.h('IC检验的是同一天的相对排序')
    g.p('如果3只股票预测分数为[0.3,0.1,−0.1]，分数排名[3,2,1]，实际收益排名[3,1,2]，无并列时Rank IC=1−6×(0²+1²+1²)/(3×(3²−1))=0.5。先对一个信号日的股票算，再对多个已成熟信号日平均。')
    g.p('普通LightGBM回归优化平方误差，模型选择又使用每日Rank IC；两者不是同一个目标。另一个ranker60才把每日股票显式分组传给排序模型。')
    g.cite('tools/research_20261011/prepare_data.py','entry=g.adj_open.shift(-1)','标签起止价格；target由同日未来收益排名产生。')
    g.cite('tools/research_20261011/train_models.py','def ic_value','评估预测与目标时重新排名；ranker60的group在fit_predict内。')

    g.page('4  时间切分：以最新2026Q4模型为例')
    q=EX[1]
    g.h('直观解释')
    g.p('用更早的数据学规律，用随后一段历史比较设置，到了季度边界再用当时已成熟的数据重拟合，预测下一个季度。因子滚动窗口、模型训练窗口、验证窗口和预测期限是不同的“时间长度”。')
    g.table(['阶段','实际信号日期 / 标签边界','作用'],[
        ('初始训练',f"{q['screening_train_first']}至{q['screening_train_last']}\n最后标签{q['screening_last_label']}",'499个有效信号日；筛选因子并训练供验证的模型。'),
        ('验证','2026-03-31至2026-09-30\n跨入测试的目标值置为缺失','比较模型、TopN及风险规则；IC只用已成熟标签。'),
        ('测试前重拟合',f"{q['refit_first']}至{q['refit_last']}\n最后标签{q['refit_last_label']}",f"{q['refit_rows']:,}行；最新代码会重新筛因子并拟合，用于下一季度。"),
        ('外层预测','2026-10-08起的Q4\n当前快照只有10月8、9日','两天各100行；不拿这两天的未来答案训练这套模型。')],[1.05,2.65,2.3],9.5)
    g.h('排除跨界标签的公式')
    g.p('初始训练样本要求 label_exit < 验证开始；上线前重拟合要求 label_exit < 测试开始。即使t早于边界，只要其未来答案晚于边界，当时就不能把它当作已知训练答案。')
    g.p('例：9月22日信号的5日标签在9月30日结束，可以进入10月8日前的重拟合。临近测试起点、需要10月行情才能知道答案的行会剔除。代码按交易日序列移动，周末和假期不当作交易日。')
    g.cite('tools/research_20261011/train_models.py',"cutoff=test.date.min()",'main()构建季度训练/验证窗口与标签边界。')
    g.cite('tools/research_20261011/train_models.py','outer_train=panel','main()取最近504个名义交易日；重新筛选后拟合部署模型。')
    g.p('这里的“上线前”描述回放实验的信息时点。目前是本地研究，不是已运行的实盘服务。旧40因子的最终重拟合用更长的约630日范围，新版用最近504日，不能混用日期解释。','Caption')

    g.page('5  每天预测：换输入，沿用本季度模型')
    g.h('直观解释与本地复核')
    g.p('fit负责学习模型内部参数；predict只把新X代入已学映射。旧模型可因当天因子变化而输出不同分数。本次加载已存Q3/Q4模型，对同一测试数据分别整季预测和逐日预测，结果最大差值均为0，也与原保存预测完全一致。')
    with (AUD/'2026Q4_daily_predictions.csv').open(encoding='utf-8-sig') as f:rows=list(csv.DictReader(f))
    g.table(['信号日期','股票','lgbm100分数','参数版本'],[[r['date'],r['stockid'],f"{float(r['lgbm100']):.6f}",'同一2026Q4模型'] for r in rows],[1.4,1.2,1.4,2],9.5)
    g.p('上表为真实保存输出。两天的5日标签在当前快照中均未成熟，但仍能预测。Q4自动选择的score实际来自xgboost60，不能把文件最后一列或score自动叫作LightGBM输出；本表专门读取lgbm100列。','Caption')
    g.h('必要公式与对应调用')
    g.p('同一季度内：sᵢ,ₜ = f_θQ(xᵢ,ₜ)，θQ固定，xᵢ,ₜ每天更新。LightGBM不接收date/stockid列，也没有隐藏的RNN记忆；mom20等输入已总结了该股过去20日的信息。')
    g.code('X = train[features].fillna(0).to_numpy()\ny = train["target_5"].to_numpy()\nmodel.fit(X, y)       # 季度边界拟合\nX_today = today[features].fillna(0).to_numpy()\nscores = model.predict(X_today)  # 每日推理')
    g.cite('tools/research_20261011/train_models.py','predictions[name]=m.predict(z)','fit_predict()实际对test整块预测；代码示例仅突出fit/predict，不是完整重跑命令。')
    g.p('历史批量predict可以包含后来日期的特征，因为每行独立代入固定模型；这本身不会让较早行看到较晚行。但特征构造和供应商数据可得时间仍须单独防泄漏。本次没有重新训练模型。','Caption')

    g.page('6  四种频率：重训、评分、新建批次、退出')
    g.h('先把当前实现说准确')
    g.table(['动作','主流程频率','时间用途'],[
        ('股票池 / 特征 / 评分','每日','收盘信息更新后，对当日最多100只评分。'),
        ('监督模型 / 入选因子身份','每季度','季度边界重新筛选、拟合；短窗实验单独按月。'),
        ('HMM参数 / 状态概率','参数按月；概率每日','不等于LightGBM每日重训。'),
        ('新建持仓 / 每批计划退出','每日尝试建批；h=5按第5日退出','现金允许时新批预算≤净值/5；受仓位、单股与容量限制。')],[1.6,1.5,2.9],9.5)
    g.fig('五日批次时间线.png','以信号日t为起点，5日批次覆盖t+1至t+5。后续批次会重叠，收益标签也会重叠。')
    g.p('例如t+1开A批、t+2开B批、t+3开C批；A在t+5收盘计划退出，B在t+6收盘计划退出。这个账户每日有交易机会；5日描述每个批次的计划期限。')
    g.p('接受“内部一致”后，本课保留当前t+5定义：IC标签和该批计划买卖端点一致。实际风险减仓、停牌或容量不足可能让真实退出偏离计划，账本应如实记录；IC也因此不等于成本后账户收益。若采用“每5天全仓换一次”，还需另加信号间隔规则，并在相同信号日检验IC；当前主账本没有这一规则。')
    g.cite('tools/research_20261011/ledger.py','due=prev+','Market.run()设置到期日；每天取上一日信号，按h控制新批预算。')
    g.p('预测20日、实际持有5日的历史探索是另一种设计。本轮基础教学和内部一致的基线不把它当作“20日预测、20日持有”来汇报。','Caption')

    g.page('7  这一课学完，应能怎样向老师汇报')
    g.h('可直接复述的回答')
    g.p('我们的样本单位是股票—日期。每天先形成一个横截面，把训练窗内多天的横截面按行拼接，用共同的LightGBM回归模型学习未来收益的截面排名。时间通过历史特征、按日期切分和标签成熟边界体现；主模型按季度重训，季度内每天换输入做预测。旧40因子表最右列属于另一条Ridge基线，是测试前用625个有效信号日联合估计的系数。当前5日账户每日建批、每批计划5日退出；我们会保持IC与所声明的交易时点一致。')
    g.h('先合上讲义，回答四个小问题')
    g.table(['自测','简短答案（先遮住右列）'],[
        ('100只股票×40因子×300天，无缺失，X多大？','30,000×40；每个测试日输入100×40，输出100个分数。'),
        ('同一股票今天分数变了，模型一定重训了吗？','不一定。同一模型接收新因子值，就会得到新分数。'),
        ('旧表GT112的β=0.012461，是哪天的IC？','它是多天联合拟合的Ridge系数；IC另列，含义与样本窗口不同。'),
        ('10月9日的5日收益还不知道，能预测吗？','能。预测只需要当日X和已训练模型；成熟后才能算该日IC。')],[3.2,2.8],10)
    g.h('下一次只讲第2课')
    g.p('下一课聚焦“原始行情怎样变成每天的股票池和因子输入”：用成交额手算Top100，用分红/拆股例子理解复权，再定位可得时间与缺失处理。不同时展开HMM、Jev或全部评价指标。')
    g.h('本课代码与来源')
    g.p('源码行号随本次哈希保存在核对资料中；真实推理核对见tools/learning_20261013/audit_project.py，6项一致性检查均通过。旧β解释另核对原始ridge.npz与旧面板，不用当前模型参数反推旧表。','Caption')
    g.link('LightGBM官方：二维特征输入、fit与predict','https://lightgbm.readthedocs.io/en/stable/pythonapi/lightgbm.LGBMRegressor.html')
    g.link('LightGBM官方：排序模型的group定义','https://lightgbm.readthedocs.io/en/stable/pythonapi/lightgbm.LGBMRanker.html')
    g.save()

def supporting():
    folder=OUT/'核对资料';folder.mkdir(exist_ok=True)
    for name in ['source_index.json','model_examples.json','inference_checks.json','old_table_answer.json','2026Q3_daily_predictions.csv','2026Q4_daily_predictions.csv','artifact_hashes.json']:
        shutil.copy2(AUD/name,folder/name)
    lines=['# 代码索引（2026-10-13）','以下行号对应当前审计哈希；更改源文件后应重新生成。']
    for row in INDEX:
        lines.append('\n## '+row['path'])
        lines.append(f"总行数：{row['lines']}；SHA256：`{row['sha256']}`")
        for f in row['functions']:lines.append(f"- `{f['name']}`：第{f['line']}行至第{f['end']}行。")
        if row['from_imports']:lines.append('导入依赖：'+', '.join(row['from_imports']))
    (folder/'代码索引.md').write_text('\n\n'.join(lines),'utf-8')
    (folder/'研究口径待改进.md').write_text('''# 研究口径待改进（本次只说明，不重写已有结果）

1. `train_models.screening`对已按日排名的特征和目标计算相关统计；缺失时使用不同边际均值/分母，并非对每个因子都先在共同有效股票上重新排名的严格成对Spearman。因子筛选分数不能直接与旧版严格每日Spearman混称完全同算法。`ic_value`对模型评估则先丢弃缺失，再按日重新排名。
2. `recent_diagnosis.py`中名为`mean_daily_rank_ic`的字段，实际直接使用评分与排名目标的Pearson相关，没有对预测再排名。旧报告08.2的该表应读作“评分与排名目标的相关”，不应继续称其为严格Rank IC。第1课示例与新核对使用`ic_value`定义。
3. `choose()`即使全部验证目标为负也会选一个相对最好的候选，没有“验证未达门槛则全现金”的显式选项。2026Q4的h5就是负验证目标仍选xgboost60；不意味着验证已达标。
4. 最新版验证前与最终重拟合前分别筛因子，身份可以变化；全部都使用当时已成熟标签，但不能声称两次使用同一固定名单。旧版40因子最终重拟合则保持原名单。
5. 现有主要运行使用5bp/边、季度模型和每日分批持仓；老师希望的1bp简化基线、月度嵌套验证和完整PIT行业风格中性化尚待新实验。已有1bp费用敏感性不是重新选型后的1bp主实验。
6. 用户已明确t+5或t+6本身不是待确认事项，以标签、信号日与计划持有端点内部一致为准。现有h5标签和批次计划退出均用t+5；风险减仓或无法成交的实际偏离另计。
7. 所有历史窗口已多次查看，训练与测试按时间隔离不能消除研究者反复试验造成的选择偏差。模型已保存不等于已完成生产部署或前瞻绩效验证。
''','utf-8')
    (OUT/'课程进度.md').write_text('''# 当前课程进度

- 全局导读：已交付。
- 第1课：样本、旧表Ridge β、时间轴、季度训练与每日预测：已交付。
- 第2课：数据与动态股票池：下一课，等待先用自己的话复述第1课。
- 第3—7课：因子筛选、模型、账户、评价、状态与扩展：按全局导读逐课讲解；本轮没有预先堆成一本长讲义。

教学顺序固定为直观解释、数字例子、公式、源码定位、自测。若回答暴露误解，先用更小例子纠正，再向后推进。
''','utf-8')
    manifest={p.relative_to(OUT).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in OUT.rglob('*') if p.is_file() and p.name!='delivery_manifest.json'}
    (OUT/'delivery_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),'utf-8')

if __name__=='__main__':
    for image in (AUD/"figures").glob("*.png"):shutil.copy2(image,ASSET/image.name)
    global_guide();first_lesson();supporting();print(OUT)
